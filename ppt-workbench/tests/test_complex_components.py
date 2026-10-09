"""Complex components bind explicit data and preserve native object isolation."""

import copy
import sys
import tempfile
import unittest
from unittest import mock
from pathlib import Path
import xml.etree.ElementTree as ET

from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts" / "svg"))
from reference_layouts.components import build_scene, catalog, compose, example
from reference_layouts.common import Svg
from drawingml.pipeline import render_svg_drawingml
from drawingml.pipeline import validate_svg_drawingml
from build_component import EXAMPLE_CONTEXT, attach_context


class ComplexComponentTests(unittest.TestCase):
    def test_memory_component_with_context_passes_existing_svg_gate(self):
        component = catalog()["mech-three-layer-agent-memory"]
        _, _, width, height = component["box"]
        svg, native, _ = compose([{"id": "component", "component_id": component["id"],
                              "x": (1280 - width) / 2, "y": (720 - height) / 2,
                              "data": example(component)}])
        svg = attach_context(svg, EXAMPLE_CONTEXT, baseline=(720 + height) / 2 + 22)
        result = validate_svg_drawingml(svg_text=svg, native_data=native, artifact_kind="component")
        self.assertTrue(result["ok"], result["issues"])

    def test_component_crop_cannot_drop_declared_native_tables(self):
        components = catalog()
        component = components["mech-impact-reversibility-matrix"]
        data = example(component)
        changed = copy.deepcopy(components)
        changed[component["id"]]["box"] = [42, 450, 1203, 264]
        with mock.patch("reference_layouts.components.catalog", return_value=changed):
            with self.assertRaisesRegex(ValueError, "native.*table"):
                build_scene(component["id"], data)

    def test_wbr_root_component_does_not_require_unselected_flow_labels(self):
        component = catalog()["mech-wbr-anomaly-root-threshold"]
        svg, native, _ = build_scene(component["id"], example(component))
        self.assertNotIn("Next Review", "".join(ET.fromstring(svg).itertext()))
        self.assertEqual(len(native["charts"]), 0)
        self.assertEqual(len(native["tables"]), 2)

    def test_five_card_values_preserve_precision_with_fixed_role_fonts(self):
        component = catalog()["rd-five-metric-result-row"]
        data = example(component)
        for slot in (0, 4):
            with self.subTest(slot=slot):
                changed = copy.deepcopy(data)
                changed["results"][slot]["value"] = "97.2%→99.1%"
                svg, _, _ = build_scene(component["id"], changed)
                nodes = [node for node in ET.fromstring(svg)
                         if node.tag.endswith("text")]
                before = next(node for node in nodes if "".join(node.itertext()) == "97.2%")
                after = next(node for node in nodes if "".join(node.itertext()) == "99.1")
                self.assertEqual(before.get("font-size"), "22")
                self.assertEqual(after.get("font-size"), "28")
                self.assertEqual(after.get("data-max-lines"), "1")
                self.assertTrue(any("".join(node.itertext()) == "%" and node.get("font-size") == "13"
                                    for node in nodes))
        data["results"][4]["value"] = "12345678901234567890%→99.1%"
        with self.assertRaisesRegex(ValueError, "exceeds fixed"):
            build_scene(component["id"], data)

    def test_fixed_five_point_radar_rejects_out_of_scale_values(self):
        component = catalog()["mech-radar-action-loop"]
        data = example(component)
        data["radar"]["series"][1]["values"][0] = 8.6
        with self.assertRaisesRegex(ValueError, "radar.*0.*5"):
            build_scene(component["id"], data)

    def test_selected_overflow_text_is_rejected_instead_of_disappearing(self):
        component = catalog()["rd-five-metric-result-row"]
        data = example(component)
        data["results"][0]["label"] = "研发端到端验证与交付综合周期指标" * 8
        with self.assertRaisesRegex(ValueError, "text exceeds.*capacity"):
            build_scene(component["id"], data)

    def test_fixed_icon_aliases_are_not_required_business_inputs(self):
        for component_id, field, aliases in [
            ("rd-three-stage-execution-roadmap", "stages", {"milestone_icon"}),
            ("rd-dual-collaboration-loops", "collaborations", {"left_icon", "right_icon"}),
        ]:
            with self.subTest(component=component_id):
                component = catalog()[component_id]
                data = example(component)
                for item in data[field]:
                    self.assertFalse(set(item) & aliases)
                svg, _, _ = build_scene(component_id, data)
                icons = [node.get("data-icon-id") for node in ET.fromstring(svg)
                         if node.get("data-icon-id")]
                self.assertTrue(icons)
                self.assertTrue(all(icons))

    def test_explicit_text_region_capacity_preserves_or_rejects_whole_text(self):
        scene = Svg()
        scene.text(20, 40, "first\nsecond", 16, width=120, max_lines=2)
        self.assertEqual("".join(scene.root.itertext()), "firstsecond")
        with self.assertRaisesRegex(ValueError, "text exceeds.*capacity"):
            scene.text(20, 40, "first\nsecond\nthird", 16, width=120, max_lines=2)

    def test_component_export_skips_only_full_page_fill_requirement(self):
        svg = '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1280 720"><g id="matrix"><rect x="50" y="200" width="1100" height="220" fill="#EFF4FA"/><text x="100" y="240" font-size="16" fill="#172438">Explicit component data</text></g></svg>'
        page = validate_svg_drawingml(svg_text=svg)
        self.assertIn("SVG_LAYOUT_VERTICAL_UNDERUSED", [issue["code"] for issue in page["issues"]])
        component = validate_svg_drawingml(svg_text=svg, artifact_kind="component")
        self.assertTrue(component["ok"], component["issues"])
        self.assertTrue(component["svg_summary"]["font_size_report"]["checked"])
        overflow = svg.replace('x="50"', 'x="1300"')
        self.assertFalse(validate_svg_drawingml(svg_text=overflow, artifact_kind="component")["ok"])

    def test_metrics_table_is_independent_from_other_page_network_data(self):
        component = catalog()["mech-causal-metrics-table"]
        svg, native, _ = build_scene(component["id"], example(component))
        self.assertEqual(len(native["tables"]), 1)
        self.assertEqual(len(native["charts"]), 0)
        self.assertNotIn("研发费用", "".join(ET.fromstring(svg).itertext()))

    def test_explicit_data_does_not_inherit_an_unselected_page_section(self):
        component = catalog()["rd-five-metric-result-row"]
        data = example(component)
        data["results_heading"] = "本次订单履约结果"
        for item in data["results"]:
            item["label"] = "本次明确提供的指标"
        svg, _, report = build_scene(component["id"], data)
        text = "".join(ET.fromstring(svg).itertext())
        self.assertIn("本次订单履约结果", text)
        self.assertIn("本次明确提供的指标", text)
        self.assertNotIn("标准模块库", text)
        self.assertEqual(report["data_fields"], component["fields"])
        with self.assertRaisesRegex(ValueError, "missing explicit fields"):
            build_scene(component["id"], {"results_heading": "缺少结果数据"})

    def test_two_case_components_keep_distinct_native_chart_workbooks(self):
        component = catalog()["rd-four-step-case-funnel"]
        first = example(component)
        second = copy.deepcopy(first)
        second["case_heading"] = "第二组明确输入的案例"
        second["case_steps"][0]["value"] = "延期12天"
        second["case_steps"][0]["chart"]["series"][0]["values"] = [3, 4, 5]
        second["case_steps"][3]["value"] = "≤8天/项"
        second["case_steps"][3]["chart"]["series"][0]["values"] = [12, 8]
        svg, native, _ = compose([
            {"id": "first", "component_id": component["id"], "x": 40, "y": 50, "data": first},
            {"id": "second", "component_id": component["id"], "x": 40, "y": 370, "data": second},
        ])
        slot_ids = [item["slot_id"] for item in native["charts"]]
        self.assertEqual(len(slot_ids), 4)
        self.assertEqual(len(set(slot_ids)), 4)
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp)
            source = path / "source.svg"
            source.write_text(svg, encoding="utf-8")
            result = render_svg_drawingml(svg_path=source, native_data=native, output=path / "page.pptx")
            self.assertTrue(result["ok"], result.get("issues"))
            charts = [shape.chart for shape in Presentation(path / "page.pptx").slides[0].shapes if shape.has_chart]
            self.assertEqual(list(charts[0].series[0].values), [9, 10, 7])
            self.assertEqual(list(charts[2].series[0].values), [3, 4, 5])

    def test_component_cannot_be_stretched_or_placed_outside_canvas(self):
        component = catalog()["rd-five-metric-result-row"]
        with self.assertRaisesRegex(ValueError, "does not fit"):
            compose([{"id": "overflow", "component_id": component["id"], "x": 200, "y": 0,
                      "data": example(component)}])


if __name__ == "__main__":
    unittest.main()
