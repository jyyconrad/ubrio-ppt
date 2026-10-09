import importlib.util
import base64
import os
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "scripts/extract_materials.py"


def load_script():
    spec = importlib.util.spec_from_file_location("extract_materials", SCRIPT)
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


class MaterialsExtractorTests(unittest.TestCase):
    def test_plain_text_scope_and_chinese_space_path(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp) / "中文 材料"
            root.mkdir()
            (root / "说明.md").write_text("第一行\n第二行\n第三行", encoding="utf-8")
            completed = subprocess.run(
                [sys.executable, str(SCRIPT), str(root), "--max-chars", "8"],
                check=False,
                capture_output=True,
                text=True,
            )
            self.assertEqual(completed.returncode, 0, completed.stderr)
            self.assertIn("来源：", completed.stdout)
            self.assertIn("[已截断]", completed.stdout)

    def test_csv_uses_real_rows_and_include_scope(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            (root / "data.csv").write_text('name,value\n"A,区",10\n', encoding="utf-8")
            (root / "ignore.md").write_text("ignore", encoding="utf-8")
            completed = subprocess.run(
                [sys.executable, str(SCRIPT), str(root), "--include", ".csv"],
                check=False,
                capture_output=True,
                text=True,
            )
            self.assertEqual(completed.returncode, 0, completed.stderr)
            self.assertIn("2: A,区 | 10", completed.stdout)
            self.assertNotIn("ignore.md", completed.stdout)

    def test_refuses_missing_input_and_silent_overwrite(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / "a.txt"
            output = root / "result.md"
            source.write_text("ok", encoding="utf-8")
            output.write_text("keep", encoding="utf-8")
            completed = subprocess.run(
                [sys.executable, str(SCRIPT), str(source), "--output", str(output)],
                check=False,
                capture_output=True,
                text=True,
            )
            self.assertEqual(completed.returncode, 2)
            self.assertEqual(output.read_text(encoding="utf-8"), "keep")
            missing = subprocess.run(
                [sys.executable, str(SCRIPT), str(root / "missing.txt")],
                check=False,
                capture_output=True,
                text=True,
            )
            self.assertEqual(missing.returncode, 2)

    def test_page_parser_rejects_invalid_range(self) -> None:
        module = load_script()
        with self.assertRaises(Exception):
            module.parse_pages("4-2")
        with self.assertRaises(Exception):
            module.parse_pages(",")
        with self.assertRaises(Exception):
            module.parse_pages("1-1000000000000")
        with self.assertRaises(Exception):
            module.parse_pages("1-10001")
        with self.assertRaises(Exception):
            module.parse_pages("100001")

    def test_empty_pages_cli_is_not_treated_as_all_pages(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "input.txt"
            path.write_text("content", encoding="utf-8")
            completed = subprocess.run(
                [sys.executable, str(SCRIPT), str(path), "--pages", ""],
                check=False,
                capture_output=True,
                text=True,
            )
            self.assertEqual(completed.returncode, 2)
            self.assertIn("页码选择不能为空", completed.stderr)

    def test_xlsx_sheet_scope_and_formula_cache_warning(self) -> None:
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "经营 数据.xlsx"
            workbook = Workbook()
            selected = workbook.active
            selected.title = "区域销售"
            selected.append(["区域", "销售额", "完成率"])
            selected.append(["华东", 120, "=B2/200"])
            other = workbook.create_sheet("内部说明")
            other["A1"] = "不应输出"
            workbook.save(path)
            completed = subprocess.run(
                [sys.executable, str(SCRIPT), str(path), "--sheet", "区域销售"],
                check=False,
                capture_output=True,
                text=True,
            )
            self.assertEqual(completed.returncode, 0, completed.stderr)
            self.assertIn("工作表 区域销售", completed.stdout)
            self.assertIn("公式缓存缺失", completed.stdout)
            self.assertNotIn("不应输出", completed.stdout)

    def test_docx_and_pptx_are_read_locally(self) -> None:
        from docx import Document
        from pptx import Presentation
        from pptx.util import Inches

        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            docx_path = root / "访谈纪要.docx"
            document = Document()
            document.add_heading("客户反馈", level=1)
            document.add_paragraph("交付周期需要缩短。")
            document.save(docx_path)

            pptx_path = root / "参考稿.pptx"
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
            box.text = "第一页结论"
            presentation.slides.add_slide(presentation.slide_layouts[6])
            presentation.save(pptx_path)

            completed = subprocess.run(
                [sys.executable, str(SCRIPT), str(docx_path), str(pptx_path), "--pages", "1"],
                check=False,
                capture_output=True,
                text=True,
            )
            self.assertEqual(completed.returncode, 0, completed.stderr)
            self.assertIn("交付周期需要缩短", completed.stdout)
            self.assertIn("第一页结论", completed.stdout)
            self.assertNotIn("幻灯片 2", completed.stdout)

    def test_out_of_range_pptx_and_pdf_pages_fail(self) -> None:
        from pptx import Presentation
        from pypdf import PdfWriter

        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            pptx_path = root / "one.pptx"
            presentation = Presentation()
            presentation.slides.add_slide(presentation.slide_layouts[6])
            presentation.save(pptx_path)
            pdf_path = root / "one.pdf"
            writer = PdfWriter()
            writer.add_blank_page(width=100, height=100)
            with pdf_path.open("wb") as handle:
                writer.write(handle)
            for path in (pptx_path, pdf_path):
                completed = subprocess.run(
                    [sys.executable, str(SCRIPT), str(path), "--pages", "2"],
                    check=False,
                    capture_output=True,
                    text=True,
                )
                self.assertEqual(completed.returncode, 1)
                self.assertIn("页码超出范围", completed.stdout)

    def test_selected_pptx_page_only_extracts_its_images_and_refuses_overwrite(self) -> None:
        from pptx import Presentation
        from pptx.util import Inches

        png = base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
        )
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            first = root / "first.png"
            second = root / "second.png"
            first.write_bytes(png)
            second.write_bytes(png)
            pptx_path = root / "images.pptx"
            presentation = Presentation()
            for image in (first, second):
                slide = presentation.slides.add_slide(presentation.slide_layouts[6])
                slide.shapes.add_picture(str(image), Inches(1), Inches(1))
            presentation.save(pptx_path)
            output = root / "extracted"
            command = [
                sys.executable,
                str(SCRIPT),
                str(pptx_path),
                "--pages",
                "1",
                "--extract-images",
                str(output),
            ]
            first_run = subprocess.run(command, check=False, capture_output=True, text=True)
            self.assertEqual(first_run.returncode, 0, first_run.stderr)
            extracted = sorted((output / "images").glob("*"))
            self.assertEqual([path.name for path in extracted], ["slide-1-image-1.png"])
            original = extracted[0].read_bytes()
            second_run = subprocess.run(command, check=False, capture_output=True, text=True)
            self.assertEqual(second_run.returncode, 1)
            self.assertIn("未写入任何文件", second_run.stdout)
            self.assertEqual(extracted[0].read_bytes(), original)

    def test_directory_discovery_skips_symlinks_outside_scope(self) -> None:
        if not hasattr(os, "symlink"):
            self.skipTest("symlink unsupported")
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            allowed = root / "allowed"
            outside = root / "outside"
            allowed.mkdir()
            outside.mkdir()
            (allowed / "inside.md").write_text("inside", encoding="utf-8")
            (outside / "secret.md").write_text("SECRET", encoding="utf-8")
            (allowed / "escape.md").symlink_to(outside / "secret.md")
            (allowed / "escape-dir").symlink_to(outside, target_is_directory=True)
            completed = subprocess.run(
                [sys.executable, str(SCRIPT), str(allowed)],
                check=False,
                capture_output=True,
                text=True,
            )
            self.assertEqual(completed.returncode, 0, completed.stderr)
            self.assertIn("inside", completed.stdout)
            self.assertNotIn("SECRET", completed.stdout)
            self.assertNotIn("secret.md", completed.stdout)

    def test_input_and_ooxml_expansion_limits_fail_before_parse(self) -> None:
        from pptx import Presentation

        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            text = root / "large.txt"
            text.write_text("123456", encoding="utf-8")
            too_large = subprocess.run(
                [sys.executable, str(SCRIPT), str(text), "--max-input-bytes", "2"],
                check=False,
                capture_output=True,
                text=True,
            )
            self.assertEqual(too_large.returncode, 2)
            pptx_path = root / "deck.pptx"
            presentation = Presentation()
            presentation.slides.add_slide(presentation.slide_layouts[6])
            presentation.save(pptx_path)
            members = subprocess.run(
                [sys.executable, str(SCRIPT), str(pptx_path), "--max-ooxml-members", "1"],
                check=False,
                capture_output=True,
                text=True,
            )
            self.assertEqual(members.returncode, 2)
            expanded = subprocess.run(
                [sys.executable, str(SCRIPT), str(pptx_path), "--max-ooxml-expanded-bytes", "1"],
                check=False,
                capture_output=True,
                text=True,
            )
            self.assertEqual(expanded.returncode, 2)

    def test_replacement_decoding_and_empty_xlsx_rows_are_reported_safely(self) -> None:
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            encoded = root / "legacy.txt"
            encoded.write_bytes(b"ok\xffbad")
            decoded = subprocess.run(
                [sys.executable, str(SCRIPT), str(encoded)],
                check=False,
                capture_output=True,
                text=True,
            )
            self.assertEqual(decoded.returncode, 0)
            self.assertIn("替换字符", decoded.stdout)
            workbook = Workbook()
            sheet = workbook.active
            sheet["C3"] = "only"
            path = root / "sparse.xlsx"
            workbook.save(path)
            sparse = subprocess.run(
                [sys.executable, str(SCRIPT), str(path)],
                check=False,
                capture_output=True,
                text=True,
            )
            self.assertEqual(sparse.returncode, 0, sparse.stderr)
            self.assertIn("范围 A1:C3", sparse.stdout)

    def test_grouped_text_and_chart_data_are_extracted(self) -> None:
        from pptx import Presentation
        from pptx.chart.data import ChartData
        from pptx.enum.chart import XL_CHART_TYPE
        from pptx.util import Inches

        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "objects.pptx"
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            group = slide.shapes.add_group_shape()
            child = group.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
            child.text = "分组文本"
            data = ChartData()
            data.categories = ["华东", "华南"]
            data.add_series("销售额", (120, 80))
            slide.shapes.add_chart(
                XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(4), Inches(1), Inches(4), Inches(3), data
            )
            presentation.save(path)
            completed = subprocess.run(
                [sys.executable, str(SCRIPT), str(path)],
                check=False,
                capture_output=True,
                text=True,
            )
            self.assertEqual(completed.returncode, 0, completed.stderr)
            self.assertIn("分组文本", completed.stdout)
            self.assertIn("华东 | 华南", completed.stdout)
            self.assertIn("销售额: 120.0 | 80.0", completed.stdout)

    def test_standalone_read_only_package_with_chinese_space_path(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            package = root / "只读 技能" / "ppt-materials"
            shutil.copytree(ROOT, package)
            source = root / "输入 材料.md"
            source.write_text("独立包运行", encoding="utf-8")
            for path in sorted(package.rglob("*"), reverse=True):
                path.chmod(0o555 if path.is_dir() else 0o444)
            package.chmod(0o555)
            completed = subprocess.run(
                [sys.executable, str(package / "scripts/extract_materials.py"), str(source)],
                check=False,
                capture_output=True,
                text=True,
            )
            self.assertEqual(completed.returncode, 0, completed.stderr)
            self.assertIn("独立包运行", completed.stdout)

    def test_original_copy_is_byte_identical(self) -> None:
        repo = ROOT.parents[1]
        source = repo / "backend/resources/agent-content/skills/file-processing"
        copied = ROOT / "references/original/file-processing"
        source_files = sorted(path.relative_to(source) for path in source.rglob("*") if path.is_file())
        copied_files = sorted(path.relative_to(copied) for path in copied.rglob("*") if path.is_file())
        self.assertEqual(source_files, copied_files)
        for relative in source_files:
            self.assertEqual((source / relative).read_bytes(), (copied / relative).read_bytes())


if __name__ == "__main__":
    unittest.main()
