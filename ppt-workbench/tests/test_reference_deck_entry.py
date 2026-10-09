"""Thin reference-deck entry: preflight the actual specs/content before rendering."""

from __future__ import annotations

import hashlib
import importlib.util
import io
import json
import sys
import tempfile
import unittest
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

from openpyxl import load_workbook
from pptx import Presentation


ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "scripts" / "svg" / "build_reference_deck.py"
SVG_SCRIPTS = ROOT / "scripts" / "svg"
SCRIPTS = ROOT / "scripts"
EXAMPLE_SPEC = ROOT / "assets" / "examples" / "blue-dense" / "result-dashboard.page_spec.json"
EXAMPLE_MANIFEST = ROOT / "assets" / "examples" / "blue-dense" / "deck-manifest.json"


def load_entry():
    for path in (str(SVG_SCRIPTS), str(SCRIPTS)):
        if path not in sys.path:
            sys.path.insert(0, path)
    spec = importlib.util.spec_from_file_location("build_reference_deck", SCRIPT)
    if spec is None or spec.loader is None:
        raise FileNotFoundError(SCRIPT)
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def _business_fixture(reference_id: str) -> dict:
    from build_reference_page import resources
    from reference_layouts.components import _business_data

    _, defaults, _ = resources()
    return _business_data(defaults[reference_id])


def _page(reference_id: str, page_no: int) -> dict:
    spec = json.loads(EXAMPLE_SPEC.read_text(encoding="utf-8"))
    spec["page_no"] = page_no
    spec["reference_id"] = reference_id
    spec["title"] = "四项指标改善，继续固化机制并验证因果"
    spec["period"] = "演示期起点与演示期当前"
    spec["population"] = "合成演示同口径样本"
    spec["denominator"] = "各指标页内口径"
    spec["limitation"] = "不代表真实企业数据；不支持因果归因"
    content = _business_fixture(reference_id)
    content["title"] = spec["title"]
    spec["reference_content"] = content
    return spec


def _document(*reference_ids: str) -> dict:
    return {
        "manifest": json.loads(EXAMPLE_MANIFEST.read_text(encoding="utf-8")),
        "pages": [_page(reference_id, index) for index, reference_id in enumerate(reference_ids, start=1)],
    }


def _all_workbook_values(pptx_path: Path) -> list[object]:
    values: list[object] = []
    with zipfile.ZipFile(pptx_path) as archive:
        workbook_names = [name for name in archive.namelist() if name.startswith("ppt/embeddings/") and name.endswith(".xlsx")]
        for name in workbook_names:
            workbook = load_workbook(io.BytesIO(archive.read(name)), data_only=True, read_only=True)
            for sheet in workbook.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    values.extend(row)
    return values


class ReferenceDeckEntryTests(unittest.TestCase):
    def test_multiple_metric_metadata_keeps_actual_values(self):
        entry = load_entry()
        spec = _page("R05", 1)
        for field in ("period", "population", "denominator", "source_ref"):
            spec.pop(field, None)
        spec["metrics"] = [
            {"period": "上半年", "population": "样本甲", "denominator": "订单", "source_ref": "台账甲"},
            {"period": "下半年", "population": "样本乙", "denominator": "项目", "source_ref": "台账乙"},
        ]

        metadata = entry._source_metadata(entry._page_spec(spec), 1)

        self.assertEqual(metadata["period"], "上半年；下半年")
        self.assertEqual(metadata["population"], "样本甲；样本乙")
        self.assertNotIn("page_spec", "".join(metadata.values()))

    def test_describe_reference_exposes_business_contract_without_fixture_facts(self):
        entry = load_entry()
        payload = entry.describe_reference("R10")

        self.assertEqual(payload["reference_id"], "R10")
        self.assertIn("chart_groups", payload["business_fields"])
        self.assertNotIn("synthetic_example", payload)
        chart_item = payload["business_data_contract"]["properties"]["chart_groups"]
        self.assertEqual(chart_item["type"], "array")
        self.assertIn("Agent", payload["agent_requirement"])

    def test_renderer_source_band_is_used_for_shared_r10_r31_content(self):
        entry = load_entry()
        described = entry.describe_reference("R10")
        self.assertEqual(described["source_band"]["font_size"], 11)
        self.assertEqual(described["source_band"]["max_lines"], 2)
        self.assertEqual(entry.describe_reference("R31")["source_band"], described["source_band"])

        prepared = entry.preflight_reference_deck(_document("R10", "R31"))
        self.assertTrue(prepared["preflight"]["ok"])
        first, second = prepared["pages"]
        self.assertEqual(first["source_footnote"]["band_origin"], "renderer_schema")
        self.assertEqual(first["source_footnote"]["band"], described["source_band"])
        self.assertEqual(first["source_footnote"]["band"], second["source_footnote"]["band"])

    def test_r27_table_footnote_leaves_room_for_evidence_source_band(self):
        entry = load_entry()
        prepared = entry.preflight_reference_deck(_document("R27"))
        self.assertTrue(prepared["preflight"]["ok"])
        self.assertEqual(prepared["pages"][0]["source_footnote"]["band_origin"], "default_bottom")

    def test_r32_uses_single_line_bottom_band_without_changing_canvas(self):
        entry = load_entry()
        prepared = entry.preflight_reference_deck(_document("R32"))
        page = prepared["pages"][0]

        self.assertEqual(page["canvas"], "1080x720")
        self.assertEqual(page["source_footnote"]["band_origin"], "default_bottom")
        self.assertEqual(page["source_footnote"]["line_count"], 1)
        self.assertEqual(page["source_footnote"]["display_state"], "观察")
        self.assertGreaterEqual(page["source_footnote"]["band"]["y"], 705)

    def test_real_deck_and_gold_validators_block_before_render(self):
        entry = load_entry()
        document = _document("R05")
        document["manifest"].pop("audience_decision")
        document["pages"][0]["metrics"] = []

        with tempfile.TemporaryDirectory() as tmp:
            output_dir = Path(tmp) / "out"
            with self.assertRaises(entry.ReferenceDeckError) as raised:
                entry.build_reference_deck(document, output_dir=output_dir)

            report = raised.exception.report
            self.assertEqual(raised.exception.code, "PREFLIGHT_BLOCKED")
            self.assertIn("manifest.audience_decision", report["deck_validation"]["blockers"])
            self.assertEqual(report["pages"][0]["gold_quality"]["quality_status"], "needs_repair")
            self.assertEqual(list(output_dir.glob("**/*.pptx")), [])

    def test_title_shape_and_second_page_capacity_fail_before_render(self):
        entry = load_entry()

        bad_title = _document("R05")
        bad_title["pages"][0]["reference_content"]["title"] = "与规格不一致"
        with self.assertRaises(entry.ReferenceDeckError) as raised:
            entry.preflight_reference_deck(bad_title)
        self.assertEqual(raised.exception.code, "TITLE_MISMATCH")

        bad_shape = _document("R05")
        bad_shape["pages"][0]["reference_content"].pop("issues")
        with self.assertRaises(entry.ReferenceDeckError) as raised:
            entry.preflight_reference_deck(bad_shape)
        self.assertEqual(raised.exception.code, "REFERENCE_CONTENT_INVALID")

        bad_capacity = _document("R05", "R07")
        bad_capacity["pages"][1]["limitation"] = "容量测试" * 900
        with tempfile.TemporaryDirectory() as tmp:
            output_dir = Path(tmp) / "out"
            with self.assertRaises(entry.ReferenceDeckError) as raised:
                entry.build_reference_deck(bad_capacity, output_dir=output_dir)
            self.assertEqual(raised.exception.code, "SOURCE_FOOTNOTE_OVERFLOW")
            self.assertEqual(list(output_dir.glob("**/*.pptx")), [])

    def test_r31_rejects_business_content_that_is_not_shared_with_r10(self):
        entry = load_entry()
        document = _document("R10", "R31")
        document["pages"][1]["reference_content"]["diagnosis"][0]["text"] = "另一套来源事实"

        with self.assertRaises(entry.ReferenceDeckError) as raised:
            entry.preflight_reference_deck(document)

        self.assertEqual(raised.exception.code, "R31_SOURCE_MISMATCH")

    def test_real_two_page_build_consumes_actual_content_and_writes_editable_pptx(self):
        entry = load_entry()
        document = _document("R05", "R07")
        chart_step = document["pages"][1]["reference_content"]["case_steps"][0]
        chart_step["value"] = "测试值97天"
        chart_step["chart"]["series"][0]["values"][-1] = 97.125
        expected_content = document["pages"][1]["reference_content"]
        expected_content_hash = hashlib.sha256(
            json.dumps(expected_content, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
        ).hexdigest()

        with tempfile.TemporaryDirectory() as tmp:
            output_dir = Path(tmp) / "out"
            manifest = entry.build_reference_deck(document, output_dir=output_dir)

            self.assertTrue(manifest["ok"])
            self.assertEqual(manifest["visual_review"], "not_performed")
            self.assertEqual(len(manifest["pages"]), 2)
            event_names = [event["event"] for event in manifest["events"]]
            self.assertLess(event_names.index("preflight_complete"), event_names.index("first_render_started"))
            self.assertEqual(event_names.count("page_gold_validated"), 2)
            self.assertEqual(event_names.count("page_svg_validated"), 2)

            for index in (1, 2):
                page_dir = output_dir / f"page-{index:03d}"
                for name in (
                    "page_spec.json",
                    "page_spec.normalized.json",
                    "reference-content.json",
                    "preflight-report.json",
                    "gold-start.svg",
                    "gold-start.native-data.json",
                    "reference-final.svg",
                    "native-data.json",
                    "adaptation-report.json",
                    "page.pptx",
                ):
                    self.assertTrue((page_dir / name).is_file(), f"missing {page_dir / name}")

            second = output_dir / "page-002"
            svg = (second / "reference-final.svg").read_text(encoding="utf-8")
            native = json.loads((second / "native-data.json").read_text(encoding="utf-8"))
            self.assertIn("测试值97天", "".join(ET.fromstring(svg).itertext()))
            self.assertEqual(native["charts"][0]["series"][0]["values"][-1], 97.125)
            self.assertIn(97.125, _all_workbook_values(second / "page.pptx"))

            page_record = manifest["pages"][1]
            self.assertEqual(page_record["sha256"]["reference_content"], expected_content_hash)
            self.assertEqual(page_record["source_footnote"]["font_px"], 11)
            self.assertLessEqual(page_record["source_footnote"]["line_count"], 2)

            with zipfile.ZipFile(second / "page.pptx") as archive:
                slide_xml = archive.read("ppt/slides/slide1.xml")
                self.assertGreater(slide_xml.count(b"<p:sp>"), 5)
                self.assertIn(b"<p:graphicFrame>", slide_xml)

            presentation = Presentation(second / "page.pptx")
            source_groups = [
                shape
                for shape in presentation.slides[0].shapes
                if hasattr(shape, "shapes")
                and any("观察｜期:" in getattr(child, "text", "") for child in shape.shapes)
            ]
            self.assertEqual(len(source_groups), 1)
            source_group = source_groups[0]
            self.assertLessEqual(source_group.left + source_group.width, presentation.slide_width)
            self.assertLessEqual(source_group.top + source_group.height, presentation.slide_height)


if __name__ == "__main__":
    unittest.main()
