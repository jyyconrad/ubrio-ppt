from __future__ import annotations

import json
import importlib.util
import os
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches


SCRIPT = Path(__file__).parents[1] / "scripts" / "merge_export.py"
PPTXGEN_FIXTURES = Path(__file__).with_name("make_pptxgenjs_fixtures.cjs")


def _load_module():
    spec = importlib.util.spec_from_file_location("merge_export_under_test", SCRIPT)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def _deck(path: Path, labels: list[str], *, chart=False, image=False, notes=False) -> None:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333333)
    presentation.slide_height = Inches(7.5)
    presentation.slides._sldIdLst.remove(presentation.slides._sldIdLst[0]) if presentation.slides else None
    for position, label in enumerate(labels):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        textbox = slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(5), Inches(0.7))
        textbox.text_frame.text = label
        table = slide.shapes.add_table(2, 2, Inches(1), Inches(1.5), Inches(3), Inches(1.2)).table
        table.cell(0, 0).text = "项目"
        table.cell(0, 1).text = "数值"
        table.cell(1, 0).text = label
        table.cell(1, 1).text = str(position + 1)
        if chart and position == 0:
            data = ChartData()
            data.categories = ["A", "B"]
            data.add_series("系列", (1, 2))
            slide.shapes.add_chart(
                XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(5), Inches(1.5), Inches(4), Inches(3), data
            )
        if image and position == 0:
            image_path = path.with_suffix(".png")
            Image.new("RGB", (80, 60), (20, 120, 200)).save(image_path)
            slide.shapes.add_picture(str(image_path), Inches(10), Inches(1.5), width=Inches(2))
        if notes and position == 0:
            slide.notes_slide.notes_text_frame.text = f"备注 {label}"
    presentation.save(path)


def _run(*args: str, env: dict[str, str] | None = None, cwd: Path | None = None) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        [sys.executable, str(SCRIPT), *args], capture_output=True, text=True, env=env, cwd=cwd
    )


def test_preflight_selects_reorders_duplicates_and_rejects_dimension_mismatch(tmp_path: Path) -> None:
    first = tmp_path / "外部 A.pptx"
    second = tmp_path / "外部 B.pptx"
    _deck(first, ["A1", "A2"])
    _deck(second, ["B1", "B2", "B3"])
    completed = _run(
        "preflight", "--input", f"{second}:3,1", "--input", f"{first}:1,2,1"
    )
    assert completed.returncode == 0, completed.stderr
    payload = json.loads(completed.stdout)
    assert payload["merge_available"] is True
    assert [(Path(item["path"]).name, item["source_page"]) for item in payload["pages"]] == [
        ("外部 B.pptx", 3),
        ("外部 B.pptx", 1),
        ("外部 A.pptx", 1),
        ("外部 A.pptx", 2),
        ("外部 A.pptx", 1),
    ]

    deck = Presentation(second)
    deck.slide_width = Inches(10)
    deck.save(second)
    mismatch = _run("preflight", "--input", str(first), "--input", str(second))
    assert mismatch.returncode == 2
    assert "dimension mismatch" in mismatch.stderr

    huge = _run("preflight", "--input", f"{first}:1-999999999")
    assert huge.returncode == 2
    assert "out of range" in huge.stderr


def test_export_requires_explicit_backend(tmp_path: Path) -> None:
    source = tmp_path / "source.pptx"
    _deck(source, ["A"])
    completed = _run("export", "--input", str(source), "--format", "pdf", "--output", str(tmp_path / "x.pdf"))
    assert completed.returncode == 2
    assert "no conversion backend selected" in completed.stderr


def test_png_overwrite_rejects_unrelated_files_and_input_directory(tmp_path: Path) -> None:
    source = tmp_path / "source.pptx"
    _deck(source, ["A"])
    output_dir = tmp_path / "pages"
    output_dir.mkdir()
    (output_dir / "keep.txt").write_text("do not delete", encoding="utf-8")
    unrelated = _run(
        "export",
        "--input",
        str(source),
        "--format",
        "png",
        "--backend",
        "libreoffice",
        "--output",
        str(output_dir),
        "--overwrite",
    )
    assert unrelated.returncode == 2
    assert "unrelated files" in unrelated.stderr
    assert (output_dir / "keep.txt").read_text(encoding="utf-8") == "do not delete"

    contains_input = _run(
        "export",
        "--input",
        str(source),
        "--format",
        "png",
        "--backend",
        "libreoffice",
        "--output",
        str(tmp_path),
        "--overwrite",
    )
    assert contains_input.returncode == 2
    assert "must not contain the input PPTX" in contains_input.stderr


@pytest.mark.skipif(
    os.environ.get("PPT_MERGE_TEST_LIBREOFFICE") != "1",
    reason="optional LibreOffice route; opt in with PPT_MERGE_TEST_LIBREOFFICE=1",
)
def test_optional_libreoffice_exports_external_deck_to_pdf_and_png(tmp_path: Path) -> None:
    source_a = tmp_path / "外部 图表 备注.pptx"
    source_b = tmp_path / "外部 图片.pptx"
    source = tmp_path / "合并结果.pptx"
    _deck(source_a, ["A1", "A2"], chart=True, notes=True)
    _deck(source_b, ["B1"], image=True)
    merged = _run(
        "merge", "--input", str(source_a), "--input", str(source_b), "--output", str(source)
    )
    assert merged.returncode == 0, merged.stderr
    pdf = tmp_path / "结果.pdf"
    completed = _run(
        "export", "--input", str(source), "--format", "pdf", "--backend", "libreoffice", "--output", str(pdf)
    )
    assert completed.returncode == 0, completed.stderr
    assert pdf.read_bytes().startswith(b"%PDF")

    png_dir = tmp_path / "预览 图片"
    completed = _run(
        "export", "--input", str(source), "--format", "png", "--backend", "libreoffice", "--output", str(png_dir)
    )
    assert completed.returncode == 0, completed.stderr
    payload = json.loads(completed.stdout)
    assert payload["pages"] == 3
    assert len(list(png_dir.glob("slide-*.png"))) == 3


def test_merge_preserves_validated_static_subset_and_refuses_overwrite(tmp_path: Path) -> None:
    first = tmp_path / "first.pptx"
    second = tmp_path / "second.pptx"
    _deck(first, ["A1", "A2"], chart=True, notes=True)
    _deck(second, ["B1", "B2", "B3"], image=True)
    target = tmp_path / "x.pptx"
    completed = _run(
        "merge",
        "--input",
        f"{second}:3,1",
        "--input",
        f"{first}:1,2,1",
        "--output",
        str(target),
    )
    assert completed.returncode == 0, completed.stderr
    merged = Presentation(target)
    titles = [
        next(shape.text for shape in slide.shapes if hasattr(shape, "text_frame") and shape.text)
        for slide in merged.slides
    ]
    assert titles == ["B3", "B1", "A1", "A2", "A1"]
    assert "备注 A1" in merged.slides[2].notes_slide.notes_text_frame.text
    assert "备注 A1" in merged.slides[4].notes_slide.notes_text_frame.text
    with zipfile.ZipFile(target) as archive:
        names = archive.namelist()
        assert len([name for name in names if name.startswith("ppt/charts/chart") and name.endswith(".xml")]) >= 2
        assert len([name for name in names if name.startswith("ppt/embeddings/") and name.endswith(".xlsx")]) >= 2
        assert any(name.startswith("ppt/media/") for name in names)
        assert len([name for name in names if name.startswith("ppt/notesSlides/notesSlide") and name.endswith(".xml")]) == 2
        assert any(name.startswith("ppt/slideMasters/slideMaster") for name in names)
        assert any(name.startswith("ppt/theme/theme") for name in names)

    existing_merge = _run(
        "merge", "--input", str(first), "--output", str(target)
    )
    assert existing_merge.returncode == 2
    assert "output exists" in existing_merge.stderr

    existing = tmp_path / "existing.pdf"
    existing.write_bytes(b"keep")
    overwrite = _run(
        "export", "--input", str(first), "--format", "pdf", "--backend", "libreoffice", "--output", str(existing)
    )
    assert overwrite.returncode == 2
    assert existing.read_bytes() == b"keep"


def test_pptxgenjs_four_source_merge_removes_orphan_root_notes_and_validates_all_relationships(
    tmp_path: Path,
) -> None:
    generated = subprocess.run(
        ["node", str(PPTXGEN_FIXTURES), str(tmp_path)],
        cwd=SCRIPT.parents[1],
        capture_output=True,
        text=True,
    )
    assert generated.returncode == 0, generated.stderr
    sources = [tmp_path / f"{name}.pptx" for name in ("03-target", "05-table", "06-process", "09-evidence")]
    output = tmp_path / "four-source-merged.pptx"
    args = ["merge"]
    for source in sources:
        args.extend(("--input", str(source)))
    args.extend(("--output", str(output)))
    completed = _run(*args)
    assert completed.returncode == 0, completed.stderr
    payload = json.loads(completed.stdout)
    assert payload["slides"] == 4
    assert isinstance(payload["removed_orphan_parts"], list)
    assert "ppt/slides/slide1.xml" in payload["removed_orphan_parts"]
    assert "ppt/notesSlides/notesSlide1.xml" in payload["removed_orphan_parts"]

    module = _load_module()
    module._validate_all_relationship_targets(output)
    merged = Presentation(output)
    assert len(merged.slides) == 4
    assert [slide.notes_slide.notes_text_frame.text.splitlines()[0] for slide in merged.slides] == [
        "Synthetic PptxGenJS note: 03-target",
        "Synthetic PptxGenJS note: 05-table",
        "Synthetic PptxGenJS note: 06-process",
        "Synthetic PptxGenJS note: 09-evidence",
    ]
    with zipfile.ZipFile(output) as archive:
        names = set(archive.namelist())
        assert "ppt/slides/slide1.xml" not in names
        assert "ppt/slides/_rels/slide1.xml.rels" not in names
        assert "ppt/notesSlides/notesSlide1.xml" not in names
        assert "ppt/notesSlides/_rels/notesSlide1.xml.rels" not in names
        assert any(name.startswith("ppt/charts/chart") and name.endswith(".xml") for name in names)
        assert any(name.startswith("ppt/embeddings/") and name.endswith(".xlsx") for name in names)
        assert any(name.startswith("ppt/media/") and not name.endswith("/") for name in names)


def test_rejects_output_equal_input_and_preserves_existing_output_on_failure(tmp_path: Path) -> None:
    source = tmp_path / "source.pptx"
    _deck(source, ["A"])
    before = source.read_bytes()
    same = _run(
        "merge", "--input", str(source), "--output", str(source), "--overwrite"
    )
    assert same.returncode == 2
    assert source.read_bytes() == before

    unsafe = tmp_path / "unsafe.pptx"
    _deck(unsafe, ["unsafe"])
    with zipfile.ZipFile(unsafe, "a") as archive:
        archive.writestr(
            "ppt/slides/_rels/unsafe.rels",
            '<?xml version="1.0" encoding="UTF-8"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rIdUnsafe" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" '
            'Target="https://example.invalid/" TargetMode="External"/></Relationships>',
        )
    existing = tmp_path / "existing.pptx"
    existing.write_bytes(b"existing-output")
    failed = _run(
        "merge", "--input", str(unsafe), "--output", str(existing), "--overwrite"
    )
    assert failed.returncode == 2
    assert "external relationship" in failed.stderr
    assert existing.read_bytes() == b"existing-output"


def test_exclusive_file_publication_rejects_racing_output(tmp_path: Path) -> None:
    module = _load_module()
    staged = tmp_path / "staged.pptx"
    output = tmp_path / "output.pptx"
    staged.write_bytes(b"new")
    output.write_bytes(b"racing-writer")
    with pytest.raises(module.ToolError, match="appeared during conversion"):
        module._publish_file(staged, output, overwrite=False)
    assert output.read_bytes() == b"racing-writer"
    assert staged.read_bytes() == b"new"


def test_relationship_target_resolution_handles_absolute_and_rejects_root_escape() -> None:
    module = _load_module()
    assert (
        module._resolve_relationship_target(
            "ppt/notesSlides/notesSlide1.xml", "/ppt/slides/slide1.xml"
        )
        == "ppt/slides/slide1.xml"
    )
    assert (
        module._resolve_relationship_target(
            "ppt/notesSlides/notesSlide1.xml", "../slides/slide1.xml"
        )
        == "ppt/slides/slide1.xml"
    )
    with pytest.raises(module.ToolError, match="escapes package root"):
        module._resolve_relationship_target("ppt/notesSlides/notesSlide1.xml", "../../../outside.xml")


def test_full_relationship_validation_rejects_sidecar_without_owner(tmp_path: Path) -> None:
    module = _load_module()
    source = tmp_path / "ownerless-rel.pptx"
    _deck(source, ["A"])
    with zipfile.ZipFile(source, "a") as archive:
        archive.writestr(
            "ppt/charts/_rels/chart999.xml.rels",
            '<?xml version="1.0" encoding="UTF-8"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/package" '
            'Target="../../presentation.xml"/></Relationships>',
        )
    with pytest.raises(module.ToolError, match="relationship part has no owner"):
        module._validate_all_relationship_targets(source)


def test_read_only_standalone_copy_merges_without_office_from_arbitrary_cwd(tmp_path: Path) -> None:
    skill_source = SCRIPT.parents[1]
    installed = tmp_path / "只读 skill"
    shutil.copytree(
        skill_source,
        installed,
        ignore=shutil.ignore_patterns("node_modules", "__pycache__", ".pytest_cache"),
    )
    install = subprocess.run(
        [
            "npm",
            "ci",
            "--offline",
            "--ignore-scripts",
            "--no-audit",
            "--no-fund",
            "--workspaces=false",
        ],
        cwd=installed,
        capture_output=True,
        text=True,
    )
    assert install.returncode == 0, install.stderr
    first = tmp_path / "输入 A.pptx"
    second = tmp_path / "输入 B.pptx"
    output = tmp_path / "输出.pptx"
    _deck(first, ["A"])
    _deck(second, ["B"])
    for path in installed.rglob("*"):
        path.chmod(0o555 if path.is_dir() else 0o444)
    installed.chmod(0o555)
    env = os.environ.copy()
    env["PATH"] = f"{Path(shutil.which('node')).parent}:/usr/bin:/bin"
    blocker = tmp_path / "block-pypdf"
    blocker.mkdir()
    (blocker / "sitecustomize.py").write_text(
        "import builtins\n"
        "_real_import = builtins.__import__\n"
        "def _blocked(name, *args, **kwargs):\n"
        "    if name == 'pypdf' or name.startswith('pypdf.'):\n"
        "        raise ImportError('pypdf intentionally unavailable')\n"
        "    return _real_import(name, *args, **kwargs)\n"
        "builtins.__import__ = _blocked\n",
        encoding="utf-8",
    )
    env["PYTHONPATH"] = str(blocker)
    command = [
        sys.executable,
        str(installed / "scripts" / "merge_export.py"),
        "merge",
        "--input",
        str(first),
        "--input",
        str(second),
        "--output",
        str(output),
    ]
    try:
        completed = subprocess.run(command, cwd=tmp_path, env=env, capture_output=True, text=True)
        assert completed.returncode == 0, completed.stderr
        assert len(Presentation(output).slides) == 2
    finally:
        installed.chmod(0o755)
        for path in installed.rglob("*"):
            path.chmod(0o755 if path.is_dir() else 0o644)
