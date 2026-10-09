"""Native tables and compact trends keep data editable in dense components."""

import copy
import sys
import tempfile
import unittest
import xml.etree.ElementTree as ET
from pathlib import Path

from pptx import Presentation
from pptx.enum.chart import XL_LABEL_POSITION, XL_TICK_LABEL_POSITION

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts" / "svg"))
from reference_layouts.common import Svg
from drawingml.pipeline import render_svg_drawingml
from overlay_native_slots import diagnose_native_slots
from build_native_chart import build_scene, validate_content
from build_reference_page import _check_content


class ReferenceSharedStyleTests(unittest.TestCase):
    def test_staggered_labels_do_not_add_labels_to_reference_series(self):
        scene = Svg()
        scene.chart("micro", 50, 200, 166, 153, ["Q1", "Q2", "Q3", "Q4"],
                    [{"name": "周期", "values": [118, 108, 99, 94]}], kind="line",
                    options={"legend": False, "target": 100, "reference_line": True,
                             "data_label_layout": "staggered"})
        svg, native = scene.finish()
        with tempfile.TemporaryDirectory() as temp:
            source, output = Path(temp) / "source.svg", Path(temp) / "page.pptx"
            source.write_text(svg, encoding="utf-8")
            result = render_svg_drawingml(svg_path=source, native_data=native,
                                         output=output, artifact_kind="component")
            self.assertTrue(result["ok"], result.get("issues"))
            chart = next(s.chart for s in Presentation(output).slides[0].shapes if s.has_chart)
            labels = chart.series[1]._element.findall(
                "{http://schemas.openxmlformats.org/drawingml/2006/chart}dLbls")
            self.assertEqual(len(labels), 1)
            self.assertEqual(labels[0].find(
                "{http://schemas.openxmlformats.org/drawingml/2006/chart}delete").get("val"), "1")
            self.assertFalse(labels[0].findall(
                "{http://schemas.openxmlformats.org/drawingml/2006/chart}dLbl"))

    def test_micro_chart_staggered_native_labels_preserve_exact_values(self):
        for kind, values, positions in [
            ("bar", [67, 67, 83, 83],
             [XL_LABEL_POSITION.INSIDE_END, XL_LABEL_POSITION.OUTSIDE_END] * 2),
            ("line", [118, 108, 99, 94],
             [XL_LABEL_POSITION.ABOVE, XL_LABEL_POSITION.BELOW] * 2),
        ]:
            with self.subTest(kind=kind):
                scene = Svg()
                scene.chart("micro", 50, 200, 166, 153,
                            ["Q1", "Q2", "Q3", "Q4"],
                            [{"name": "观察", "values": values}], kind=kind,
                            fmt='0"%"', options={"legend": False,
                                "data_label_layout": "staggered", "data_label_wrap": False})
                svg, native = scene.finish()
                with tempfile.TemporaryDirectory() as temp:
                    source, output = Path(temp) / "source.svg", Path(temp) / "page.pptx"
                    source.write_text(svg, encoding="utf-8")
                    result = render_svg_drawingml(svg_path=source, native_data=native,
                                                 output=output, artifact_kind="component")
                    self.assertTrue(result["ok"], result.get("issues"))
                    chart = next(s.chart for s in Presentation(output).slides[0].shapes if s.has_chart)
                    self.assertEqual(list(chart.series[0].values), values)
                    self.assertEqual([p.data_label.position for p in chart.series[0].points], positions)
                    self.assertTrue(all(p.data_label.font.size.pt == 9
                                        for p in chart.series[0].points))
                    for point in chart.series[0].points:
                        label = point.data_label._dLbl
                        number_format = label.find("{http://schemas.openxmlformats.org/drawingml/2006/chart}numFmt")
                        self.assertIsNotNone(number_format)
                        self.assertEqual(number_format.get("formatCode"), '0"%"')
                        body = label.find("{http://schemas.openxmlformats.org/drawingml/2006/chart}txPr/"
                                          "{http://schemas.openxmlformats.org/drawingml/2006/main}bodyPr")
                        self.assertEqual(body.get("wrap"), "none")

    def test_external_radar_labels_keep_native_categories_and_scale(self):
        scene = Svg()
        scene.chart("radar", 50, 200, 480, 210, ["选择", "价格", "便利", "速度", "信任"],
                    [{"name": "当前", "values": [3.9, 3.4, 4.2, 3.6, 4.0]}],
                    kind="radar", options={"category_labels": False, "legend": False,
                        "value_axis": {"minimum": 0, "maximum": 5, "labels": False}})
        svg, native = scene.finish()
        with tempfile.TemporaryDirectory() as temp:
            source, output = Path(temp) / "source.svg", Path(temp) / "page.pptx"
            source.write_text(svg, encoding="utf-8")
            result = render_svg_drawingml(svg_path=source, native_data=native,
                                         output=output, artifact_kind="component")
            self.assertTrue(result["ok"], result.get("issues"))
            chart = next(shape.chart for shape in Presentation(output).slides[0].shapes if shape.has_chart)
            self.assertEqual(chart.category_axis.tick_label_position, XL_TICK_LABEL_POSITION.NONE)
            self.assertEqual(chart.value_axis.tick_label_position, XL_TICK_LABEL_POSITION.NONE)
            self.assertTrue(chart.category_axis.visible)
            self.assertEqual(chart.value_axis.maximum_scale, 5)
            self.assertEqual(list(chart.series[0].values), [3.9, 3.4, 4.2, 3.6, 4.0])
            self.assertEqual([item.label for item in chart.plots[0].categories],
                             ["选择", "价格", "便利", "速度", "信任"])

    def test_reference_numeric_content_rejects_nonfinite_values(self):
        for value in (float("nan"), float("inf"), float("-inf")):
            with self.subTest(value=value):
                with self.assertRaisesRegex(ValueError, "finite"):
                    _check_content({"score": value}, {"score": 1})

    def test_headerless_table_keeps_all_body_rows_and_body_style(self):
        scene = Svg()
        scene.table("causes", 50, 200, 800, 150, ["层级", "证据"],
                    [["接口", "15项说明缺失"], ["验证", "8项记录缺失"]],
                    show_header=False)
        svg, native = scene.finish()
        with tempfile.TemporaryDirectory() as temp:
            output = Path(temp) / "page.pptx"
            source = Path(temp) / "source.svg"
            source.write_text(svg, encoding="utf-8")
            result = render_svg_drawingml(svg_path=source, native_data=native,
                                         output=output, artifact_kind="component")
            self.assertTrue(result["ok"], result.get("issues"))
            table = next(shape.table for shape in Presentation(output).slides[0].shapes
                         if shape.has_table)
            self.assertEqual(len(table.rows), 2)
            self.assertEqual(table.cell(0, 0).text, "接口")
            self.assertFalse(table.first_row)
            self.assertEqual(str(table.cell(0, 0).fill.fore_color.rgb), "EFF4FA")
            self.assertFalse(table.cell(0, 0).text_frame.paragraphs[0].runs[0].font.bold)

    def test_native_sparkline_uses_exact_data_without_axes_or_small_slot_warning(self):
        scene = Svg()
        scene.chart("weekly", 50, 200, 190, 60,
                    ["W-5", "W-4", "W-3", "W-2", "W-1", "W0"],
                    [{"name": "交付", "values": [0.72, 0.75, 0.78, 0.77, 0.8, 0.82]}],
                    kind="line", fmt="0%", options={"sparkline": True})
        svg, native = scene.finish()
        slot = next(node for node in ET.fromstring(svg) if node.get("id") == "weekly")
        self.assertEqual(slot.get("data-native-chart-variant"), "sparkline")
        report = diagnose_native_slots(svg, native)
        self.assertFalse(report["issues"], report["issues"])
        self.assertNotIn("NATIVE_CHART_SLOT_TOO_SMALL", [item["code"] for item in report["warnings"]])
        with tempfile.TemporaryDirectory() as temp:
            output = Path(temp) / "page.pptx"
            source = Path(temp) / "source.svg"
            source.write_text(svg, encoding="utf-8")
            result = render_svg_drawingml(svg_path=source, native_data=native,
                                         output=output, artifact_kind="component")
            self.assertTrue(result["ok"], result.get("issues"))
            self.assertNotIn("NATIVE_CHART_SLOT_TOO_SMALL", [item["code"] for item in result["warnings"]])
            chart = next(shape.chart for shape in Presentation(output).slides[0].shapes
                         if shape.has_chart)
            self.assertEqual(list(chart.series[0].values), [0.72, 0.75, 0.78, 0.77, 0.8, 0.82])
            self.assertEqual([item.label for item in chart.plots[0].categories],
                             ["W-5", "W-4", "W-3", "W-2", "W-1", "W0"])
            self.assertFalse(chart.category_axis.visible)
            self.assertFalse(chart.value_axis.visible)
            self.assertFalse(chart.has_legend)
            self.assertTrue(chart.plots[0].has_data_labels)
            self.assertGreaterEqual(chart.plots[0].data_labels.font.size.pt, 9)

    def test_sparkline_cannot_hide_full_chart_requirements(self):
        scene = Svg()
        scene.chart("bad", 50, 200, 190, 60, ["一", "二"],
                    [{"name": "A", "values": [1, 2]}], kind="line",
                    options={"sparkline": True})
        svg, native = scene.finish()
        for patch in ({"kind": "bar"}, {"series": [native["charts"][0]["series"][0]] * 2},
                      {"options": {"sparkline": False}},
                      {"options": {"sparkline": True, "reference_line": True, "target": 2}}):
            with self.subTest(patch=patch):
                changed = copy.deepcopy(native)
                changed["charts"][0].update(patch)
                report = diagnose_native_slots(svg, changed)
                self.assertIn("NATIVE_SPARKLINE_INVALID", [item["code"] for item in report["issues"]])

    def test_quick_reference_caption_does_not_repeat_series_label_or_percent_unit(self):
        raw = {"title": "准时率变化", "unit": "百分比", "period": "2026年7至9月",
               "source": "授权合成履约台账", "limitation": "合成演示，不代表实绩",
               "labels": ["七月", "八月", "九月"],
               "series": [{"name": "实际", "values": [0.972, 0.985, 0.991]}],
               "target": {"label": "目标99.5%", "value": 0.995}, "number_format": "0.0%"}
        content = validate_content("line", raw)
        svg, native = build_scene("line", content)
        text = "\n".join("".join(node.itertext()) for node in ET.fromstring(svg)
                         if node.tag.endswith("text"))
        self.assertIn("参考值：99.5%", text)
        self.assertNotIn("目标99.5% 99.5%", text)
        self.assertIn("单位：百分比 | 期间：2026年7至9月", text)
        self.assertEqual(native["charts"][0]["options"]["reference_label"], "目标99.5%")


if __name__ == "__main__":
    unittest.main()
