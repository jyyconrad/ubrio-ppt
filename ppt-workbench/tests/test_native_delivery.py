import copy
import importlib.util
import unittest
import xml.etree.ElementTree as ET
from pathlib import Path

from pptx import Presentation
from pptx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]


def load_module(name, relative):
    definition = importlib.util.spec_from_file_location(name, ROOT / relative)
    module = importlib.util.module_from_spec(definition)
    definition.loader.exec_module(module)
    return module


class NativeDeliveryTests(unittest.TestCase):
    def setUp(self):
        self.builder = load_module('native_delivery_builder', 'scripts/svg/build_gold_svg_page.py')
        self.spec = {
            'layout': 'gold_evidence_table', 'page_role': 'content',
            'title': '先补齐交接记录，再比较推进方式',
            'primary_claim': '先补齐交接记录，再比较推进方式',
            'key_message': '先补齐交接记录，再比较推进方式',
            'audience_takeaway': '团队先核对同一任务的记录。',
            'plain_language_who_scene_problem_action': {
                'who': '团队', 'scene': '讨论交接', 'problem': '记录未齐', 'action': '核对同一任务的记录'
            },
            'metrics': [], 'quality_targets': {'metrics_min': 0},
            'typography': {'body_px': 16, 'title_px': 32, 'source_px': 11, 'font_family': 'Noto Sans CJK SC'},
            'sources': ['合成流程示例'],
            'native_table_slots': [{
                'id': 'ledger', 'box': {'x': 70, 'y': 220, 'w': 1139, 'h': 350},
                'columns': ['环节', '记录', '条件', '责任', '状态'],
                'rows': [['环节一', '记录一', '齐全', '负责人', '未执行'],
                         ['环节二', '记录二', '齐全', '负责人', '未执行'],
                         ['环节三', '记录三', '齐全', '负责人', '未执行'],
                         ['环节四', '记录四', '齐全', '负责人', '未执行']]
            }]
        }
        self.spec['columns'] = self.spec['native_table_slots'][0]['columns']
        self.spec['rows'] = copy.deepcopy(self.spec['native_table_slots'][0]['rows'])

    def test_native_table_sources_do_not_count_as_tiny_body(self):
        normalized = self.builder.normalize_page_spec(self.spec)
        svg = self.builder.build_svg_page(normalized)
        result = self.builder.run_svg_gates(svg, normalized)
        self.assertNotIn('tiny_text_ratio', result.missing)
        self.assertEqual(result.details['source_text_node_count'], 2)

    def test_native_table_inherits_declared_font_and_size(self):
        self.spec['typography']['body_px'] = 22
        style = self.builder.build_native_data(self.spec)['tables'][0]['style']
        self.assertEqual(style.get('font_face'), 'Noto Sans CJK SC')
        self.assertEqual(style.get('font_size'), 16.5)

    def test_native_table_keeps_explicit_slot_style(self):
        override = {'font_face': 'Source Han Sans SC', 'font_size': 14, 'header_fill': '#123A63'}
        self.spec['native_table_slots'][0]['style'] = override
        self.assertEqual(self.builder.build_native_data(self.spec)['tables'][0]['style'], override)

    def test_fixed_page_honors_declared_font(self):
        fixed = load_module('fixed_delivery_builder', 'scripts/svg/build_fixed_svg_page.py')
        spec = {'kind': 'closing', 'theme_id': 'business_blue', 'title': '先补齐交接记录',
                'subtitle': '核对样本后再讨论', 'typography': {'font_family': 'Noto Sans CJK SC'}}
        root = ET.fromstring(fixed.build_fixed_svg_page(spec))
        texts = [node for node in root.iter() if node.tag.endswith('text')]
        self.assertTrue(texts)
        self.assertTrue(all(node.get('font-family') == 'Noto Sans CJK SC' for node in texts))

    def test_native_table_rejects_style_below_body_floor(self):
        self.spec['density_budget'] = {'min_body_font_px': 15}
        self.spec['native_table_slots'][0]['style'] = {'font_size': 10}
        normalized = self.builder.normalize_page_spec(self.spec)
        report = self.builder.quality_report(normalized)
        self.assertIn('native_table_font_px', report['missing'])

    def test_native_table_honors_higher_declared_floor(self):
        self.spec['density_budget'] = {'min_body_font_px': 20}
        self.spec['native_table_slots'][0]['style'] = {'font_size': 12}
        result = self.builder.run_page_gates(self.builder.normalize_page_spec(self.spec))
        self.assertIn('native_table_font_px', result.missing)
        self.spec['native_table_slots'][0]['style']['font_size'] = 15
        result = self.builder.run_page_gates(self.builder.normalize_page_spec(self.spec))
        self.assertNotIn('native_table_font_px', result.missing)

    def test_native_table_keeps_legal_font_override(self):
        self.spec['native_table_slots'][0]['style'] = {'font_size': 12}
        result = self.builder.run_page_gates(self.builder.normalize_page_spec(self.spec))
        self.assertNotIn('native_table_font_px', result.missing)
        self.assertEqual(self.builder.build_native_data(self.spec)['tables'][0]['style']['font_size'], 12)

    def test_native_table_sets_and_updates_east_asian_font(self):
        overlay = load_module('native_delivery_overlay', 'scripts/svg/overlay_native_slots.py')
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        cell = slide.shapes.add_table(1, 1, 0, 0, 914400, 914400).table.cell(0, 0)
        cell.text = '核对记录'
        for family in ['Noto Sans CJK SC', 'Source Han Sans SC']:
            overlay._style_table_cell(cell, row_index=1, style={'font_face': family, 'font_size': 14})
            properties = cell.text_frame.paragraphs[0].runs[0]._r.get_or_add_rPr()
            east_asian = properties.findall(qn('a:ea'))
            self.assertEqual(len(east_asian), 1)
            self.assertEqual(east_asian[0].get('typeface'), family)

    def test_too_small_source_is_still_rejected(self):
        svg = '<svg xmlns="http://www.w3.org/2000/svg"><g id="source-note"><text font-size="8">来源</text></g></svg>'
        self.assertIn('svg_body_font_px', self.builder.run_svg_gates(svg).missing)

    def test_small_body_is_not_hidden_by_source_exceptions(self):
        svg = '<svg xmlns="http://www.w3.org/2000/svg"><text font-size="11">正文</text><g id="source-note"><text font-size="11">来源</text></g></svg>'
        result = self.builder.run_svg_gates(svg)
        self.assertIn('tiny_text_ratio', result.missing)
        self.assertIn('svg_body_font_px', result.missing)


if __name__ == '__main__':
    unittest.main()
