"""The quick chart entry keeps business data exact in native PPT charts."""

from __future__ import annotations

import importlib.util
import io
import json
import math
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE

ROOT = Path(__file__).resolve().parents[1]
BUILDER = ROOT / "scripts" / "svg" / "build_native_chart.py"
PYTHON = ROOT / ".venv" / "bin" / "python"
NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
    "s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
}

COMMON = {
    "title": "交付质量对比",
    "period": "本期",
    "source": "授权合成台账",
    "limitation": "仅反映已纳入台账的样本",
}
FIXTURES = {
    "radar": {
        **COMMON,
        "unit": "分",
        "labels": ["完整性", "及时性", "一致性", "可追溯", "可复用"],
        "series": [
            {"name": "基准", "values": [4.8, 4.6, 4.7, 4.5, 4.4]},
            {"name": "当前", "values": [3.75, 4.1, 3.6, 3.9, 4.05]},
        ],
        "scale": {"minimum": 0, "maximum": 5},
        "number_format": "0.00",
    },
    "line": {
        **COMMON,
        "unit": "%",
        "labels": ["第1期", "第2期", "第3期", "第4期"],
        "series": [
            {"name": "实际", "values": [0.67, 0.725, 0.8, 0.835]},
            {"name": "计划", "values": [0.62, 0.7, 0.76, 0.82]},
            {"name": "基准", "values": [0.7, 0.74, 0.78, 0.81]},
            {"name": "上期", "values": [0.64, 0.68, 0.73, 0.79]},
        ],
        "target": {"value": 0.9, "label": "目标"},
        "number_format": "0.0%",
    },
    "bar": {
        **COMMON,
        "unit": "万元",
        "labels": ["区域甲", "区域乙", "区域丙"],
        "series": [
            {"name": "差额", "values": [12.5, -3.25, 7]},
        ],
        "target": {"value": 10, "label": "目标"},
    },
}


def _python() -> str:
    return str(PYTHON) if PYTHON.is_file() else sys.executable


def _load_builder():
    if not BUILDER.is_file():
        raise FileNotFoundError(f"missing quick chart entry: {BUILDER}")
    spec = importlib.util.spec_from_file_location("build_native_chart", BUILDER)
    if spec is None or spec.loader is None:
        raise FileNotFoundError(BUILDER)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def _chart_tags(path: Path) -> set[str]:
    tags: set[str] = set()
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            if name.startswith("ppt/charts/chart") and name.endswith(".xml"):
                root = ET.fromstring(archive.read(name))
                tags.update(element.tag.rsplit("}", 1)[-1] for element in root.iter())
    return tags


def _workbook_cells(path: Path) -> dict[str, str | float]:
    with zipfile.ZipFile(path) as package:
        workbooks = sorted(
            name for name in package.namelist()
            if name.startswith("ppt/embeddings/") and name.endswith(".xlsx")
        )
        if len(workbooks) != 1:
            raise AssertionError(f"expected one embedded workbook, found {workbooks}")
        workbook_bytes = package.read(workbooks[0])
    with zipfile.ZipFile(io.BytesIO(workbook_bytes)) as workbook:
        shared: list[str] = []
        if "xl/sharedStrings.xml" in workbook.namelist():
            root = ET.fromstring(workbook.read("xl/sharedStrings.xml"))
            shared = [
                "".join(node.text or "" for node in item.findall(".//s:t", NS))
                for item in root.findall("s:si", NS)
            ]
        sheets = sorted(
            name for name in workbook.namelist()
            if name.startswith("xl/worksheets/sheet") and name.endswith(".xml")
        )
        root = ET.fromstring(workbook.read(sheets[0]))
        cells: dict[str, str | float] = {}
        for cell in root.findall(".//s:c", NS):
            ref = cell.get("r")
            if not ref:
                continue
            cell_type = cell.get("t")
            if cell_type == "inlineStr":
                cells[ref] = "".join(
                    node.text or "" for node in cell.findall(".//s:t", NS)
                )
                continue
            raw = cell.findtext("s:v", namespaces=NS)
            if raw is None:
                continue
            if cell_type == "s":
                cells[ref] = shared[int(raw)]
            elif cell_type in {"str", "e"}:
                cells[ref] = raw
            else:
                cells[ref] = float(raw)
        return cells


def _assert_numbers_equal(test: unittest.TestCase, actual, expected) -> None:
    test.assertEqual(len(actual), len(expected))
    for got, wanted in zip(actual, expected):
        test.assertTrue(math.isclose(float(got), float(wanted), rel_tol=0, abs_tol=1e-12))


class NativeChartQuickTests(unittest.TestCase):
    def test_describe_exposes_business_contract_without_example_facts(self) -> None:
        proc = subprocess.run(
            [_python(), str(BUILDER), "--describe"],
            check=False,
            capture_output=True,
            text=True,
        )
        self.assertEqual(proc.returncode, 0, proc.stderr)
        description = json.loads(proc.stdout)
        self.assertEqual(description["supported_kinds"], ["radar", "line", "bar"])
        self.assertEqual(description["artifact_kind"], "component")
        self.assertNotIn("synthetic_example", description)
        contract_text = json.dumps(description["data_contract"], ensure_ascii=False)
        for field in ("title", "labels", "series", "unit", "period", "source", "limitation"):
            self.assertIn(field, contract_text)
        self.assertNotIn("coordinates", contract_text)
        self.assertNotIn("colors", contract_text)
        self.assertNotIn("font_size", contract_text)

    def test_radar_line_and_bar_are_native_and_keep_exact_workbook_data(self) -> None:
        expected_tags = {"radar": "radarChart", "line": "lineChart", "bar": "barChart"}
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            for kind, fixture in FIXTURES.items():
                with self.subTest(kind=kind):
                    page = root / kind
                    content = root / f"{kind}.json"
                    content.write_text(
                        json.dumps(fixture, ensure_ascii=False), encoding="utf-8"
                    )
                    proc = subprocess.run(
                        [
                            _python(),
                            str(BUILDER),
                            "--kind",
                            kind,
                            "--content",
                            str(content),
                            "--output-dir",
                            str(page),
                        ],
                        check=False,
                        capture_output=True,
                        text=True,
                    )
                    self.assertEqual(proc.returncode, 0, proc.stderr)
                    for name in (
                        "content.json",
                        "source.svg",
                        "native-data.json",
                        "page.pptx",
                        "native-chart-report.json",
                    ):
                        self.assertTrue((page / name).is_file(), name)
                    source_svg = (page / "source.svg").read_text(encoding="utf-8")
                    self.assertIn(f"单位：{fixture['unit']}", source_svg)
                    pptx = page / "page.pptx"
                    self.assertIn(expected_tags[kind], _chart_tags(pptx))
                    presentation = Presentation(pptx)
                    slide = presentation.slides[0]
                    charts = [
                        shape.chart
                        for shape in slide.shapes
                        if shape.has_chart
                    ]
                    self.assertEqual(len(charts), 1)
                    chart = charts[0]
                    target = fixture.get("target")
                    if target:
                        self.assertTrue(chart.has_legend)
                        expected_caption = {
                            "line": "参考值：90%",
                            "bar": "参考值：10 万元",
                        }[kind]
                        slide_text = "\n".join(
                            shape.text for shape in slide.shapes if hasattr(shape, "text")
                        )
                        self.assertIn(expected_caption, slide_text)
                    expected_names = [item["name"] for item in fixture["series"]]
                    expected_values = [item["values"] for item in fixture["series"]]
                    if target:
                        expected_names.append(target["label"])
                        expected_values.append([target["value"]] * len(fixture["labels"]))
                    self.assertEqual([item.name for item in chart.series], expected_names)
                    for actual, expected in zip(chart.series, expected_values):
                        _assert_numbers_equal(self, list(actual.values), expected)
                    if kind == "radar":
                        self.assertEqual(chart.chart_type, XL_CHART_TYPE.RADAR_FILLED)
                        self.assertFalse(chart.plots[0].has_data_labels)
                    if kind == "line":
                        self.assertEqual(
                            chart.value_axis.tick_labels.number_format, "0.0%"
                        )
                        colors = {
                            str(item.format.line.color.rgb)
                            for item in list(chart.series)[:4]
                        }
                        self.assertEqual(len(colors), 4)
                    if kind == "bar":
                        self.assertLessEqual(chart.value_axis.minimum_scale, -3.25)
                        self.assertEqual(
                            chart.value_axis.tick_labels.number_format, "0.##"
                        )
                    workbook = _workbook_cells(pptx)
                    for index, label in enumerate(fixture["labels"], start=2):
                        self.assertEqual(workbook[f"A{index}"], label)
                    for series_index, (name, values) in enumerate(
                        zip(expected_names, expected_values), start=2
                    ):
                        column = chr(64 + series_index)
                        self.assertEqual(workbook[f"{column}1"], name)
                        for row, value in enumerate(values, start=2):
                            self.assertTrue(
                                math.isclose(
                                    float(workbook[f"{column}{row}"]),
                                    float(value),
                                    rel_tol=0,
                                    abs_tol=1e-12,
                                )
                            )
                    report = json.loads(
                        (page / "native-chart-report.json").read_text(encoding="utf-8")
                    )
                    self.assertEqual(report["status"], "structurally_verified")
                    self.assertEqual(report["native_object_status"]["chart_count"], 1)
                    self.assertTrue(report["native_object_status"]["embedded_workbook"])
                    self.assertEqual(report["visual_review"], "not_performed")
                    self.assertEqual(report["powerpoint_roundtrip"], "not_performed")
                    warning_codes = {
                        warning["code"] for warning in report["render"].get("warnings", [])
                    }
                    self.assertNotIn("SVG_TEXT_FONT_SIZE_BELOW_MINIMUM", warning_codes)

    def test_radar_fill_target_series_and_bar_zero_baseline_use_shared_engine(self) -> None:
        builder = _load_builder()
        radar = builder.validate_content("radar", FIXTURES["radar"])
        _, radar_native = builder.build_scene("radar", radar)
        radar_chart = radar_native["charts"][0]
        self.assertFalse(radar_chart["options"]["data_labels"])
        self.assertEqual(radar_chart["options"]["radar_fill_series_index"], 1)
        self.assertEqual(radar_chart["options"]["value_axis"]["minimum"], 0)
        self.assertEqual(radar_chart["options"]["value_axis"]["maximum"], 5)
        bar = builder.validate_content("bar", FIXTURES["bar"])
        _, bar_native = builder.build_scene("bar", bar)
        bar_chart = bar_native["charts"][0]
        self.assertLess(bar_chart["options"]["value_axis"]["minimum"], -3.25)
        self.assertEqual(bar_chart["options"]["value_axis"]["format"], "0.##")
        line = builder.validate_content("line", FIXTURES["line"])
        _, line_native = builder.build_scene("line", line)
        self.assertFalse(line_native["charts"][0]["options"]["data_labels"])
        single_line = json.loads(json.dumps(FIXTURES["line"], ensure_ascii=False))
        single_line["series"] = single_line["series"][:1]
        single_line = builder.validate_content("line", single_line)
        _, single_line_native = builder.build_scene("line", single_line)
        self.assertTrue(single_line_native["charts"][0]["options"]["data_labels"])
        boundary = json.loads(json.dumps(FIXTURES["bar"], ensure_ascii=False))
        boundary["title"] = "题" * 32
        boundary["limitation"] = "限" * 78
        boundary = builder.validate_content("bar", boundary)
        boundary_svg, _ = builder.build_scene("bar", boundary)
        self.assertIn("题" * 32, boundary_svg)
        self.assertIn("限" * 78, boundary_svg)
        for kind in ("line", "bar"):
            normalized = builder.validate_content(kind, FIXTURES[kind])
            _, native = builder.build_scene(kind, normalized)
            chart = native["charts"][0]
            self.assertTrue(chart["options"]["legend"])
            self.assertTrue(chart["options"]["reference_line"])
            self.assertEqual(chart["options"]["reference_label"], "目标")

    def test_invalid_shapes_and_numbers_are_rejected_without_coercion(self) -> None:
        builder = _load_builder()
        cases = []
        length = json.loads(json.dumps(FIXTURES["line"], ensure_ascii=False))
        length["series"][0]["values"] = [1, 2]
        cases.append(("line", length, "does not match labels length"))
        nan = json.loads(json.dumps(FIXTURES["line"], ensure_ascii=False))
        nan["series"][0]["values"][0] = float("nan")
        cases.append(("line", nan, "must be finite"))
        boolean = json.loads(json.dumps(FIXTURES["bar"], ensure_ascii=False))
        boolean["series"][0]["values"][0] = True
        cases.append(("bar", boolean, "bool is not accepted"))
        missing = json.loads(json.dumps(FIXTURES["bar"], ensure_ascii=False))
        missing["series"][0]["values"][0] = None
        cases.append(("bar", missing, "missing values are not supported"))
        scale = json.loads(json.dumps(FIXTURES["radar"], ensure_ascii=False))
        scale["series"][1]["values"][0] = 5.1
        cases.append(("radar", scale, "outside scale 0..5"))
        unknown = json.loads(json.dumps(FIXTURES["line"], ensure_ascii=False))
        unknown["x"] = 64
        cases.append(("line", unknown, "unknown fields"))
        nested_unknown = json.loads(json.dumps(FIXTURES["line"], ensure_ascii=False))
        nested_unknown["series"][0]["color"] = "#000000"
        cases.append(("line", nested_unknown, "unknown fields"))
        for kind, payload, message in cases:
            with self.subTest(kind=kind, message=message):
                with self.assertRaisesRegex(builder.ContentError, message):
                    builder.validate_content(kind, payload)

    def test_text_capacity_error_is_structured_without_traceback(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            payload = json.loads(json.dumps(FIXTURES["bar"], ensure_ascii=False))
            payload["title"] = "超" * 33
            content = root / "content.json"
            content.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
            proc = subprocess.run(
                [
                    _python(),
                    str(BUILDER),
                    "--kind",
                    "bar",
                    "--content",
                    str(content),
                    "--output-dir",
                    str(root / "output"),
                ],
                check=False,
                capture_output=True,
                text=True,
            )
            self.assertEqual(proc.returncode, 2)
            self.assertNotIn("Traceback", proc.stderr)
            error = json.loads(proc.stderr)
            self.assertIn("32-character capacity", error["error"])


if __name__ == "__main__":
    unittest.main()
