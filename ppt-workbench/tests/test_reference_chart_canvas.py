"""Reference charts must remain native and align on nonstandard canvases."""

import sys
import tempfile
import unittest
from pathlib import Path

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts" / "svg"))
import overlay_native_slots as overlay
from drawingml.pipeline import render_svg_drawingml


class ReferenceChartCanvasTests(unittest.TestCase):
    def test_negative_bar_points_keep_sign_and_labels_outside_the_plot(self):
        from reference_layouts.common import Svg
        from pptx.enum.chart import XL_TICK_LABEL_POSITION
        svg = Svg()
        svg.chart("b", 50, 100, 500, 300, ["A", "B", "C"],
                  [{"name": "Actual", "values": [12.5, -3.25, 7]}],
                  options={"target": 10, "reference_line": True,
                           "value_axis": {"minimum": -3.25}})
        _, meta = svg.finish()
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        overlay._add_chart(slide, {"id": "b", "x": "50", "y": "100", "width": "500", "height": "300"}, meta, (1280, 720))
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "negative.pptx"
            presentation.save(path)
            chart = Presentation(path).slides[0].shapes[-1].chart
            self.assertEqual(list(chart.series[0].values), [12.5, -3.25, 7])
            point = chart.series[0].points[1].format.element
            invert = point.find(qn("c:invertIfNegative"))
            self.assertIsNotNone(invert)
            self.assertEqual(invert.get("val"), "0")
            self.assertEqual(chart.category_axis.tick_label_position, XL_TICK_LABEL_POSITION.LOW)

    def test_native_data_labels_can_keep_number_and_unit_on_one_line(self):
        from reference_layouts.common import Svg
        svg = Svg()
        svg.chart("m", 100, 100, 200, 90, ["W1", "W2", "W3"],
                  [{"name": "Actual", "values": [52, 58, 64]}], kind="line", fmt='0"%"',
                  options={"axes": False, "data_label_position": "below", "data_label_wrap": False})
        _, meta = svg.finish()
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        overlay._add_chart(slide, {"id": "m", "x": "100", "y": "100", "width": "200", "height": "90"}, meta, (1280, 720))
        chart = slide.shapes[-1].chart
        body = chart.plots[0]._element.find("./" + qn("c:dLbls") + "/" + qn("c:txPr") + "/" + qn("a:bodyPr"))
        self.assertIsNotNone(body)
        self.assertEqual(body.get("wrap"), "none")
        self.assertEqual(list(chart.series[0].values), [52, 58, 64])
        self.assertEqual(chart.plots[0].data_labels.number_format, '0"%"')

    def test_radar_can_fill_only_current_series_without_changing_values(self):
        from reference_layouts.common import Svg
        from pptx.enum.chart import XL_CHART_TYPE
        from pptx.enum.dml import MSO_LINE_DASH_STYLE
        svg = Svg()
        values = [[4.6, 4.5, 4.7, 4.4, 4.3], [3.8, 3.4, 4.1, 3.2, 3.6], [4.1, 3.9, 4.2, 3.8, 3.9]]
        names = ["目标", "本期", "试点"]
        svg.chart("r", 40, 100, 500, 300, ["易用", "响应", "可靠", "透明", "关怀"],
                  [{"name": name, "values": row} for name, row in zip(names, values)], kind="radar",
                  options={"data_labels": False, "radar_fill_series_index": 1,
                           "radar_fill_transparency": 65, "dashed_series_indices": [2]})
        _, meta = svg.finish()
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        attrs = {"id": "r", "x": "40", "y": "100", "width": "500", "height": "300"}
        self.assertTrue(overlay._add_chart(slide, attrs, meta, (1280, 720)))
        chart = slide.shapes[-1].chart
        self.assertEqual(chart.chart_type, XL_CHART_TYPE.RADAR_FILLED)
        self.assertFalse(chart.plots[0].has_data_labels)
        for index, series in enumerate(chart.series):
            style = series._element.find(qn("c:spPr"))
            if index == 1:
                alpha = style.find("./" + qn("a:solidFill") + "/" + qn("a:srgbClr") + "/" + qn("a:alpha"))
                self.assertEqual(alpha.get("val"), "35000")
            else:
                self.assertIsNotNone(style.find(qn("a:noFill")))
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "radar.pptx"
            presentation.save(path)
            loaded = Presentation(path).slides[0].shapes[-1].chart
            self.assertEqual([item.name for item in loaded.series], names)
            self.assertEqual([list(item.values) for item in loaded.series], values)
            self.assertEqual(loaded.series[2].format.line.dash_style, MSO_LINE_DASH_STYLE.DASH)
            self.assertNotEqual(loaded.series[0].format.line.dash_style, MSO_LINE_DASH_STYLE.DASH)

    def test_dashed_series_indices_must_reference_existing_series(self):
        attrs = {"id": "r", "x": "40", "y": "100", "width": "500", "height": "300"}
        for indices in ([True], [-1], [1], "0"):
            with self.subTest(indices=indices):
                meta = {"charts": [{"slot_id": "r", "kind": "radar", "labels": ["A", "B", "C"],
                                    "series": [{"name": "current", "values": [1, 2, 3]}],
                                    "options": {"dashed_series_indices": indices}}]}
                presentation = Presentation()
                slide = presentation.slides.add_slide(presentation.slide_layouts[6])
                with self.assertRaisesRegex(ValueError, "dashed_series_indices"):
                    overlay._add_chart(slide, attrs, meta, (1280, 720))

    def test_thresholds_and_target_category_remain_native_chart_data(self):
        from reference_layouts.common import Svg
        svg = Svg()
        svg.chart("b", 50, 100, 400, 240, ["Q1", "Q2"],
                  [{"name": "Actual", "values": [67, 83]}], fmt='0"%"',
                  options={"target": 90, "target_category": "TARGET", "reference_line": True})
        _, meta = svg.finish()
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        self.assertTrue(overlay._add_chart(slide, {"id": "b", "x": "50", "y": "100", "width": "400", "height": "240"}, meta, (1280, 720)))
        chart = slide.shapes[-1].chart
        self.assertEqual(list(chart.series[0].values), [67, 83, 90])
        self.assertEqual(list(chart.series[1].values), [90, 90, 90])
        self.assertEqual(len(chart.plots), 2)
        self.assertIsNotNone(chart._chartSpace.find('.//' + qn('c:lineChart')))
        self.assertEqual(chart.value_axis.tick_labels.number_format, '0"%"')

    def test_weekly_table_can_use_a_thin_header_and_chart_labels_below(self):
        from reference_layouts.common import Svg
        svg = Svg()
        svg.table("t", 30, 120, 600, 260, ["Goal", "Input"],
                  [["A", "B"], ["C", "D"]], row_heights=[0.25, 1, 1])
        svg.chart("c", 300, 160, 200, 100, ["W1", "W2"],
                  [{"name": "Actual", "values": [67, 72]}], kind="line",
                  options={"data_label_position": "below", "axes": False})
        _, meta = svg.finish()
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        self.assertTrue(overlay._add_table(slide, {"id": "t", "x": "30", "y": "120", "width": "600", "height": "260"}, meta, (1280, 720)))
        table = slide.shapes[-1].table
        self.assertAlmostEqual(table.rows[0].height / table.rows[1].height, 0.25, places=5)
        self.assertTrue(overlay._add_chart(slide, {"id": "c", "x": "300", "y": "160", "width": "200", "height": "100"}, meta, (1280, 720)))
        chart = slide.shapes[-1].chart
        from pptx.enum.chart import XL_LABEL_POSITION
        self.assertEqual(chart.plots[0].data_labels.position, XL_LABEL_POSITION.BELOW)

    def test_table_type_column_keeps_navigation_color_in_body_rows(self):
        from reference_layouts.common import Svg, BLUE
        svg = Svg()
        svg.table("t", 30, 400, 900, 180, ["Type", "Definition"],
                  [["Input", "Actionable"], ["Outcome", "Result"]],
                  column_styles=[{"body_fill": BLUE, "text_color": "#FFFFFF", "bold": True}])
        _, meta = svg.finish()
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        self.assertTrue(overlay._add_table(slide, {"id": "t", "x": "30", "y": "400", "width": "900", "height": "180"}, meta, (1280, 720)))
        table = slide.shapes[-1].table
        self.assertEqual(str(table.cell(1, 0).fill.fore_color.rgb), BLUE.lstrip("#"))
        self.assertEqual(str(table.cell(1, 0).text_frame.paragraphs[0].runs[0].font.color.rgb), "FFFFFF")
        self.assertTrue(table.cell(1, 0).text_frame.paragraphs[0].runs[0].font.bold)
        self.assertNotEqual(str(table.cell(1, 1).fill.fore_color.rgb), BLUE.lstrip("#"))

    def test_result_bar_keeps_zero_baseline_and_only_final_point_emphasized(self):
        from reference_layouts.common import Svg, BLUE
        svg = Svg()
        svg.chart("b", 80, 120, 400, 220, ["Q1", "Q2", "Q3", "Q4"],
                  [{"name": "Rate", "values": [67, 72, 78, 83]}])
        _, meta = svg.finish()
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        self.assertTrue(overlay._add_chart(slide, {"id": "b", "x": "80", "y": "120", "width": "400", "height": "220"}, meta, (1280, 720)))
        chart = slide.shapes[-1].chart
        self.assertEqual(chart.value_axis.minimum_scale, 0)
        colors = [str(point.format.fill.fore_color.rgb) for point in chart.series[0].points]
        self.assertEqual(colors[:3], ["A8BEDF"] * 3)
        self.assertEqual(colors[3], BLUE.lstrip("#"))

    def test_auto_canvas_preserves_three_by_two_and_cropped_two_by_one(self):
        for width, height in [(1080, 720), (1280, 640)]:
            with self.subTest(canvas=(width, height)), tempfile.TemporaryDirectory() as temp:
                root = Path(temp)
                source = root / "source.svg"
                source.write_text(f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}"><text x="40" y="80" fill="#172438" font-size="28">Canvas test</text></svg>', encoding="utf-8")
                report = render_svg_drawingml(svg_path=source, output=root / "page.pptx", canvas_format="auto")
                self.assertTrue(report["ok"], report.get("issues"))
                presentation = Presentation(root / "page.pptx")
                self.assertAlmostEqual(presentation.slide_width / presentation.slide_height, width / height, places=5)

    def test_native_radar_retains_all_three_series_and_five_dimensions(self):
        svg = '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1280 720"><rect id="r" data-role="native-chart-slot" x="640" y="200" width="500" height="300"/></svg>'
        meta = {"charts": [{"slot_id": "r", "kind": "radar",
                            "labels": ["选择", "价格", "便利", "速度", "信任"],
                            "series": [{"name": name, "values": [4, 3, 4, 3, 5]}
                                       for name in ["理想", "当前", "样本方案"]],
                            "style": {"font_face": "Noto Sans CJK SC"}}]}
        self.assertTrue(overlay.diagnose_native_slots(svg, meta)["ok"])
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        self.assertTrue(overlay._add_chart(slide, {"id": "r", "x": "640", "y": "200", "width": "500", "height": "300"}, meta, (1280, 720)))
        chart = slide.shapes[-1].chart
        self.assertEqual(len(chart.series), 3)
        self.assertIsNotNone(chart._chartSpace.find('.//' + qn('c:radarChart')))
        self.assertEqual(len(chart.plots[0].categories), 5)

    def test_native_table_uses_actual_three_by_two_slide_geometry(self):
        presentation = Presentation()
        presentation.slide_width = Inches(12)
        presentation.slide_height = Inches(8)
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        attrs = {"id": "t", "x": "540", "y": "360", "width": "270", "height": "180"}
        meta = {"tables": [{"slot_id": "t", "columns": ["字段", "值"], "rows": [["A", "B"]]}]}
        self.assertTrue(overlay._add_table(slide, attrs, meta, (1080, 720)))
        shape = slide.shapes[-1]
        self.assertAlmostEqual(shape.left.inches, 6, places=3)
        self.assertAlmostEqual(shape.top.inches, 4, places=3)
        self.assertAlmostEqual(shape.width.inches, 3, places=3)
        self.assertAlmostEqual(shape.height.inches, 2, places=3)


if __name__ == "__main__":
    unittest.main()
