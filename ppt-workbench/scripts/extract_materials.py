#!/usr/bin/env python3
"""Extract readable, source-located content from common local material files."""

from __future__ import annotations

import argparse
import csv
import io
import sys
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable


SUPPORTED = {".txt", ".md", ".markdown", ".csv", ".xlsx", ".pptx", ".docx", ".pdf"}
OOXML = {".xlsx", ".pptx", ".docx"}
MAX_PAGE_NUMBER = 100_000
MAX_SELECTED_PAGES = 10_000


@dataclass
class Extraction:
    path: Path
    sections: list[tuple[str, str]] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    images: list[str] = field(default_factory=list)
    error: str | None = None


def parse_pages(value: str | None) -> set[int] | None:
    if value is None:
        return None
    if not value.strip():
        raise argparse.ArgumentTypeError("页码选择不能为空")
    pages: set[int] = set()
    for part in value.split(","):
        part = part.strip()
        if not part:
            continue
        if "-" in part:
            start_text, end_text = part.split("-", 1)
            start, end = int(start_text), int(end_text)
            if start < 1 or end < start:
                raise argparse.ArgumentTypeError(f"无效页码范围: {part}")
            if end > MAX_PAGE_NUMBER:
                raise argparse.ArgumentTypeError(
                    f"页码不能超过 {MAX_PAGE_NUMBER}: {end}"
                )
            range_length = end - start + 1
            if range_length > MAX_SELECTED_PAGES:
                raise argparse.ArgumentTypeError(
                    f"单个页码范围最多选择 {MAX_SELECTED_PAGES} 页: {part}"
                )
            for page in range(start, end + 1):
                pages.add(page)
                if len(pages) > MAX_SELECTED_PAGES:
                    raise argparse.ArgumentTypeError(
                        f"页码选择总数不能超过 {MAX_SELECTED_PAGES}"
                    )
        else:
            page = int(part)
            if page < 1:
                raise argparse.ArgumentTypeError("页码必须从 1 开始")
            if page > MAX_PAGE_NUMBER:
                raise argparse.ArgumentTypeError(f"页码不能超过 {MAX_PAGE_NUMBER}: {page}")
            pages.add(page)
            if len(pages) > MAX_SELECTED_PAGES:
                raise argparse.ArgumentTypeError(
                    f"页码选择总数不能超过 {MAX_SELECTED_PAGES}"
                )
    if not pages:
        raise argparse.ArgumentTypeError("页码选择不能为空")
    return pages


def discover(paths: Iterable[Path], recursive: bool, include: set[str]) -> list[Path]:
    found: list[Path] = []
    for path in paths:
        if not path.exists():
            raise FileNotFoundError(f"输入不存在: {path}")
        if path.is_symlink():
            raise ValueError(f"不接受符号链接输入: {path}")
        if path.is_file():
            if path.suffix.lower() not in SUPPORTED:
                raise ValueError(f"不支持的文件类型: {path}")
            if not include or path.suffix.lower() in include:
                found.append(path)
            continue
        iterator = path.rglob("*") if recursive else path.glob("*")
        for item in iterator:
            if item.is_symlink() or not item.is_file():
                continue
            if item.suffix.lower() in SUPPORTED and (not include or item.suffix.lower() in include):
                found.append(item)
    return sorted(dict.fromkeys(item.resolve() for item in found))


def validate_input(path: Path, max_input_bytes: int, max_members: int, max_expanded_bytes: int) -> None:
    size = path.stat().st_size
    if size > max_input_bytes:
        raise ValueError(f"输入文件超过 {max_input_bytes} 字节限制: {path} ({size} 字节)")
    if path.suffix.lower() not in OOXML:
        return
    try:
        with zipfile.ZipFile(path) as archive:
            members = archive.infolist()
            if len(members) > max_members:
                raise ValueError(f"OOXML 成员数超过 {max_members} 限制: {path}")
            expanded = sum(member.file_size for member in members)
            if expanded > max_expanded_bytes:
                raise ValueError(
                    f"OOXML 展开大小超过 {max_expanded_bytes} 字节限制: {path} ({expanded} 字节)"
                )
    except zipfile.BadZipFile as exc:
        raise ValueError(f"OOXML 文件损坏或不是有效 ZIP: {path}") from exc


def validate_selected_pages(pages: set[int] | None, total: int, kind: str) -> None:
    if not pages:
        return
    invalid = sorted(page for page in pages if page > total)
    if invalid:
        rendered = ", ".join(str(page) for page in invalid)
        raise ValueError(f"{kind} 页码超出范围 1-{total}: {rendered}")


def decode_utf8(path: Path, result: Extraction) -> str:
    text = path.read_bytes().decode("utf-8-sig", errors="replace")
    replacements = text.count("\ufffd")
    if replacements:
        result.warnings.append(f"UTF-8 解码产生 {replacements} 个替换字符，源文件编码可能不是 UTF-8。")
    return text


def truncate(text: str, limit: int, result: Extraction) -> str:
    if len(text) <= limit:
        return text
    result.warnings.append(f"文本超过每文件 {limit} 字符上限，输出已截断。")
    return text[:limit] + "\n\n[已截断]"


def extract_text(path: Path, result: Extraction, limit: int) -> None:
    text = decode_utf8(path, result)
    line_count = text.count("\n") + (1 if text else 0)
    result.sections.append((f"行 1-{line_count}", truncate(text, limit, result)))


def extract_csv(path: Path, result: Extraction, limit: int) -> None:
    lines: list[str] = []
    with io.StringIO(decode_utf8(path, result), newline="") as handle:
        for row_number, row in enumerate(csv.reader(handle), start=1):
            lines.append(f"{row_number}: " + " | ".join(row))
    result.sections.append((f"CSV 行 1-{len(lines)}", truncate("\n".join(lines), limit, result)))


def extract_xlsx(path: Path, result: Extraction, limit: int, sheets: set[str]) -> None:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("读取 XLSX 需要安装 openpyxl") from exc
    values_book = load_workbook(path, read_only=True, data_only=True)
    formula_book = load_workbook(path, read_only=True, data_only=False)
    try:
        selected = sheets or set(values_book.sheetnames)
        missing = selected.difference(values_book.sheetnames)
        if missing:
            raise ValueError(f"工作表不存在: {', '.join(sorted(missing))}")
        remaining = limit
        for name in values_book.sheetnames:
            if name not in selected:
                continue
            values_sheet = values_book[name]
            formula_sheet = formula_book[name]
            rows: list[str] = []
            formula_cache_missing: list[str] = []
            max_column = 0
            max_row = 0
            for row_number, (value_row, formula_row) in enumerate(
                zip(values_sheet.iter_rows(), formula_sheet.iter_rows()), start=1
            ):
                rendered: list[str] = []
                for column_number, (value_cell, formula_cell) in enumerate(
                    zip(value_row, formula_row), start=1
                ):
                    value = value_cell.value
                    formula = formula_cell.value
                    if value is not None or formula is not None:
                        max_column = max(max_column, column_number)
                        max_row = row_number
                    if isinstance(formula, str) and formula.startswith("=") and value is None:
                        coordinate = getattr(formula_cell, "coordinate", f"R{row_number}C{column_number}")
                        formula_cache_missing.append(coordinate)
                        value = f"[公式缓存缺失: {formula}]"
                    rendered.append("" if value is None else str(value))
                if any(rendered):
                    rows.append(f"{row_number}: " + " | ".join(rendered))
            text = "\n".join(rows)
            clipped = truncate(text, max(0, remaining), result)
            remaining -= min(len(text), remaining)
            if max_column:
                from openpyxl.utils import get_column_letter

                end_coordinate = f"{get_column_letter(max_column)}{max_row}"
            else:
                end_coordinate = "A1"
            result.sections.append((f"工作表 {name}，范围 A1:{end_coordinate}", clipped))
            if formula_cache_missing:
                preview = ", ".join(formula_cache_missing[:20])
                suffix = "..." if len(formula_cache_missing) > 20 else ""
                result.warnings.append(f"工作表 {name} 的公式缓存为空: {preview}{suffix}")
    finally:
        values_book.close()
        formula_book.close()


def pptx_shape_blocks(shape, prefix: str, result: Extraction) -> list[str]:
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    blocks: list[str] = []
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for child_number, child in enumerate(shape.shapes, start=1):
            blocks.extend(pptx_shape_blocks(child, f"{prefix}.{child_number}", result))
        return blocks
    if getattr(shape, "has_text_frame", False) and shape.text.strip():
        blocks.append(f"形状 {prefix}: {shape.text.strip()}")
    if getattr(shape, "has_table", False):
        for row_number, row in enumerate(shape.table.rows, start=1):
            blocks.append(f"表格 {prefix} 行 {row_number}: " + " | ".join(cell.text for cell in row.cells))
    if getattr(shape, "has_chart", False):
        try:
            chart = shape.chart
            categories = [str(category) for category in chart.plots[0].categories]
            blocks.append(f"图表 {prefix} 类别: " + " | ".join(categories))
            for series_number, series in enumerate(chart.series, start=1):
                blocks.append(
                    f"图表 {prefix} 系列 {series_number} {series.name}: "
                    + " | ".join(str(value) for value in series.values)
                )
        except Exception as exc:
            result.warnings.append(f"图表 {prefix} 数据未能完整提取: {type(exc).__name__}: {exc}")
    return blocks


def extract_pptx(path: Path, result: Extraction, limit: int, pages: set[int] | None) -> None:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise RuntimeError("读取 PPTX 需要安装 python-pptx") from exc
    presentation = Presentation(path)
    validate_selected_pages(pages, len(presentation.slides), "PPTX")
    remaining = limit
    for page_number, slide in enumerate(presentation.slides, start=1):
        if pages and page_number not in pages:
            continue
        blocks: list[str] = []
        for shape_number, shape in enumerate(slide.shapes, start=1):
            blocks.extend(pptx_shape_blocks(shape, str(shape_number), result))
        try:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                blocks.append(f"备注: {notes}")
        except (AttributeError, KeyError):
            pass
        text = "\n".join(blocks)
        result.sections.append((f"幻灯片 {page_number}", truncate(text, max(0, remaining), result)))
        remaining -= min(len(text), remaining)


def extract_docx(path: Path, result: Extraction, limit: int) -> None:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("读取 DOCX 需要安装 python-docx") from exc
    document = Document(path)
    blocks: list[str] = []
    for number, paragraph in enumerate(document.paragraphs, start=1):
        if paragraph.text.strip():
            style = paragraph.style.name if paragraph.style else ""
            blocks.append(f"段落 {number} [{style}]: {paragraph.text.strip()}")
    for table_number, table in enumerate(document.tables, start=1):
        for row_number, row in enumerate(table.rows, start=1):
            blocks.append(f"表格 {table_number} 行 {row_number}: " + " | ".join(cell.text for cell in row.cells))
    result.sections.append(("DOCX 段落与表格（无可靠最终页码）", truncate("\n".join(blocks), limit, result)))


def extract_pdf(path: Path, result: Extraction, limit: int, pages: set[int] | None) -> None:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError("读取 PDF 需要安装 pypdf") from exc
    reader = PdfReader(path)
    if reader.is_encrypted:
        raise ValueError("PDF 已加密，未尝试解密")
    validate_selected_pages(pages, len(reader.pages), "PDF")
    remaining = limit
    for page_number, page in enumerate(reader.pages, start=1):
        if pages and page_number not in pages:
            continue
        text = page.extract_text() or ""
        if len(text.strip()) < 20:
            result.warnings.append(f"PDF 第 {page_number} 页几乎无可提取文字，可能是扫描件或字体编码不可读。")
        result.sections.append((f"PDF 第 {page_number} 页", truncate(text, max(0, remaining), result)))
        remaining -= min(len(text), remaining)


def pptx_selected_images(path: Path, pages: set[int] | None) -> list[tuple[str, bytes]]:
    from pptx import Presentation
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT

    presentation = Presentation(path)
    validate_selected_pages(pages, len(presentation.slides), "PPTX")
    images: list[tuple[str, bytes]] = []
    selected = pages or set(range(1, len(presentation.slides) + 1))
    for page_number, slide in enumerate(presentation.slides, start=1):
        if page_number not in selected:
            continue
        image_number = 0
        for relationship in slide.part.rels.values():
            if relationship.reltype != RT.IMAGE:
                continue
            image_number += 1
            suffix = Path(relationship.target_ref).suffix or ".bin"
            images.append((f"slide-{page_number}-image-{image_number}{suffix}", relationship.target_part.blob))
    return images


def extract_images(path: Path, destination: Path, result: Extraction, pages: set[int] | None) -> None:
    suffix = path.suffix.lower()
    if suffix not in {".docx", ".pptx"}:
        if suffix == ".pdf":
            result.warnings.append("当前 PDF 路线不提取图片；需要时使用独立 PDF 渲染或图像工具。")
        return
    output_dir = destination / path.stem
    if suffix == ".pptx":
        candidates = pptx_selected_images(path, pages)
    else:
        candidates = []
        with zipfile.ZipFile(path) as archive:
            for member in archive.namelist():
                if member.startswith("word/media/") and not member.endswith("/"):
                    candidates.append((Path(member).name, archive.read(member)))
    targets = [(output_dir / name, content) for name, content in candidates]
    existing = [target for target, _ in targets if target.exists()]
    if existing:
        raise FileExistsError(f"图片输出已存在，未写入任何文件: {existing[0]}")
    output_dir.mkdir(parents=True, exist_ok=True)
    for target, content in targets:
        target.write_bytes(content)
        result.images.append(str(target))


def extract(path: Path, args: argparse.Namespace) -> Extraction:
    result = Extraction(path=path)
    try:
        suffix = path.suffix.lower()
        if suffix in {".txt", ".md", ".markdown"}:
            extract_text(path, result, args.max_chars)
        elif suffix == ".csv":
            extract_csv(path, result, args.max_chars)
        elif suffix == ".xlsx":
            extract_xlsx(path, result, args.max_chars, set(args.sheet or []))
        elif suffix == ".pptx":
            extract_pptx(path, result, args.max_chars, args.pages)
        elif suffix == ".docx":
            extract_docx(path, result, args.max_chars)
        elif suffix == ".pdf":
            extract_pdf(path, result, args.max_chars, args.pages)
        if args.extract_images:
            extract_images(path, args.extract_images, result, args.pages)
    except Exception as exc:  # File-level errors should not hide other requested files.
        result.error = f"{type(exc).__name__}: {exc}"
    return result


def render(results: list[Extraction]) -> str:
    lines = ["# 素材提取结果", ""]
    for result in results:
        lines.extend([f"## {result.path.name}", "", f"来源：`{result.path}`", ""])
        if result.error:
            lines.extend([f"错误：{result.error}", ""])
            continue
        for title, text in result.sections:
            lines.extend([f"### {title}", "", text or "[无可提取文字]", ""])
        if result.images:
            lines.extend(["### 提取图片", ""] + [f"- `{item}`" for item in result.images] + [""])
        if result.warnings:
            lines.extend(["### 警告", ""] + [f"- {warning}" for warning in result.warnings] + [""])
    return "\n".join(lines).rstrip() + "\n"


def parser() -> argparse.ArgumentParser:
    command = argparse.ArgumentParser(description=__doc__)
    command.add_argument("paths", nargs="+", type=Path, help="文件或目录")
    command.add_argument("--output", type=Path, help="Markdown 输出文件；默认打印到 stdout")
    command.add_argument("--overwrite", action="store_true", help="明确允许覆盖输出文件")
    command.add_argument("--no-recursive", action="store_true", help="目录只读取第一层")
    command.add_argument("--include", help="逗号分隔扩展名，例如 .pdf,.xlsx")
    command.add_argument("--max-chars", type=int, default=100_000, help="每个文件最多输出字符数")
    command.add_argument("--max-input-bytes", type=int, default=100_000_000, help="单个输入文件字节上限")
    command.add_argument("--max-ooxml-members", type=int, default=10_000, help="OOXML ZIP 成员数上限")
    command.add_argument(
        "--max-ooxml-expanded-bytes", type=int, default=500_000_000, help="OOXML ZIP 展开总字节上限"
    )
    command.add_argument("--sheet", action="append", help="只读取指定 XLSX 工作表，可重复")
    command.add_argument("--pages", type=parse_pages, help="PDF/PPTX 页码，例如 1,3-5")
    command.add_argument("--extract-images", type=Path, help="DOCX/PPTX 图片输出目录")
    return command


def main(argv: list[str] | None = None) -> int:
    args = parser().parse_args(argv)
    if min(
        args.max_chars, args.max_input_bytes, args.max_ooxml_members, args.max_ooxml_expanded_bytes
    ) < 1:
        print("所有大小和数量限制必须大于 0", file=sys.stderr)
        return 2
    include = set()
    if args.include:
        include = {item.strip().lower() for item in args.include.split(",") if item.strip()}
        include = {item if item.startswith(".") else f".{item}" for item in include}
        unsupported = include.difference(SUPPORTED)
        if unsupported:
            print(f"--include 包含不支持类型: {', '.join(sorted(unsupported))}", file=sys.stderr)
            return 2
    try:
        files = discover(args.paths, not args.no_recursive, include)
        if not files:
            raise ValueError("指定范围内没有受支持文件")
        for path in files:
            validate_input(
                path, args.max_input_bytes, args.max_ooxml_members, args.max_ooxml_expanded_bytes
            )
        if args.output and args.output.exists() and not args.overwrite:
            raise FileExistsError(f"输出已存在，拒绝覆盖: {args.output}")
        results = [extract(path, args) for path in files]
        text = render(results)
        if args.output:
            args.output.parent.mkdir(parents=True, exist_ok=True)
            args.output.write_text(text, encoding="utf-8")
        else:
            sys.stdout.write(text)
        return 1 if any(item.error for item in results) else 0
    except (FileNotFoundError, FileExistsError, ValueError) as exc:
        print(str(exc), file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
